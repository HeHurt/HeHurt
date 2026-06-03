# %% [markdown]
## Experimental benchmark: 25C 0.25P

Load cleaned cycle data extracted from `软包-循环模版25℃-0.25P.xlsm` and compare experimental capacity retention with the CW391/CW511 simulation curves.

# %%
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

EXP_DATA_DIR = Path(r"D:\Users\hez\Desktop\hithium\data_processed\314Ah\CW391_CW511_25C_0p25P")
EXP_CYCLES_PATH = EXP_DATA_DIR / "cw391_cw511_25c_0p25p_cycles_long.csv"
EXP_GROUP_SUMMARY_PATH = EXP_DATA_DIR / "cw391_cw511_25c_0p25p_group_summary.csv"
EXP_OVERALL_SUMMARY_PATH = EXP_DATA_DIR / "cw391_cw511_25c_0p25p_overall_summary.csv"

exp_cycles = pd.read_csv(EXP_CYCLES_PATH)
exp_group_summary = pd.read_csv(EXP_GROUP_SUMMARY_PATH)
exp_overall_summary = pd.read_csv(EXP_OVERALL_SUMMARY_PATH)

print(f"Loaded {len(exp_cycles)} cycle rows from {exp_cycles['barcode'].nunique()} cells")
print("Cycle range:", int(exp_cycles["cycle"].min()), "to", int(exp_cycles["cycle"].max()))
display(exp_cycles.groupby(["group", "barcode"], as_index=False)["cycle"].max())

# %%
EXP_COLORS = {
    "A": "#4C78A8",
    "B": "#F58518",
    "D": "#54A24B",
    "G": "#B279A2",
}

def _plot_sim_capacity_retention(ax):
    if {"eq_cycle", "ret_391", "ret_511"}.issubset(globals()):
        ax.plot(eq_cycle, ret_391, label="CW391 Sim", color="black", linewidth=2.2)
        ax.plot(eq_cycle, ret_511, label="CW511 Sim", color="#666666", linewidth=2.2)
        return True
    return False

fig, ax = plt.subplots(figsize=(8, 5))

for group, frame in exp_group_summary.groupby("group"):
    color = EXP_COLORS.get(group)
    ax.plot(
        frame["cycle"],
        frame["capacity_retention_pct_mean"],
        ls="--",
        marker="o",
        markersize=3,
        linewidth=1.4,
        color=color,
        label=f"Group {group} Exp",
    )

ax.plot(
    exp_overall_summary["cycle"],
    exp_overall_summary["capacity_retention_pct_mean"],
    ls="--",
    color="#D62728",
    linewidth=2.0,
    label="Overall Exp mean",
)

has_sim = _plot_sim_capacity_retention(ax)
ax.set_xlabel("Cycle number")
ax.set_ylabel("Capacity retention (%)")
ax.set_title("Capacity retention benchmark at 25℃ 0.25P")
ax.grid(alpha=0.3)
ax.legend(ncol=2 if has_sim else 1)
fig.tight_layout()
plt.show()

# %%
fig, axes = plt.subplots(1, 2, figsize=(11, 4.2), sharex=True)

axes[0].plot(
    exp_overall_summary["cycle"],
    exp_overall_summary["energy_retention_pct_mean"],
    ls="--",
    color="#4C78A8",
    linewidth=2,
    label="Energy retention Exp mean",
)
axes[0].set_ylabel("Energy retention (%)")
axes[0].set_xlabel("Cycle number")
axes[0].set_title("Energy retention")
axes[0].grid(alpha=0.3)
axes[0].legend()

axes[1].plot(
    exp_overall_summary["cycle"],
    exp_overall_summary["energy_efficiency_pct_mean"],
    ls="--",
    color="#F58518",
    linewidth=2,
    label="Energy efficiency Exp mean",
)
axes[1].set_ylabel("Energy efficiency (%)")
axes[1].set_xlabel("Cycle number")
axes[1].set_title("Energy efficiency")
axes[1].grid(alpha=0.3)
axes[1].legend()

fig.tight_layout()
plt.show()

# %%
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

EXPORT_XLSX_PATH = EXP_DATA_DIR / "cw391_cw511_25c_0p25p_benchmark_results.xlsx"

if "exp_cycles" not in globals():
    exp_cycles = pd.read_csv(EXP_CYCLES_PATH)
if "exp_group_summary" not in globals():
    exp_group_summary = pd.read_csv(EXP_GROUP_SUMMARY_PATH)
if "exp_overall_summary" not in globals():
    exp_overall_summary = pd.read_csv(EXP_OVERALL_SUMMARY_PATH)

if {"eq_cycle", "cap_391", "cap_511", "ret_391", "ret_511"}.issubset(globals()):
    sim_export_n = min(len(eq_cycle), len(cap_391), len(cap_511), len(ret_391), len(ret_511))
    sim_capacity_retention = pd.DataFrame({
        "cycle": np.asarray(eq_cycle[:sim_export_n], dtype=int),
        "cw391_capacity_Ah": np.asarray(cap_391[:sim_export_n], dtype=float),
        "cw391_capacity_retention_pct": np.asarray(ret_391[:sim_export_n], dtype=float),
        "cw511_capacity_Ah": np.asarray(cap_511[:sim_export_n], dtype=float),
        "cw511_capacity_retention_pct": np.asarray(ret_511[:sim_export_n], dtype=float),
    })
else:
    sim_capacity_retention = pd.DataFrame(columns=[
        "cycle",
        "cw391_capacity_Ah",
        "cw391_capacity_retention_pct",
        "cw511_capacity_Ah",
        "cw511_capacity_retention_pct",
    ])

benchmark_capacity = exp_overall_summary[[
    "cycle",
    "n_cells",
    "discharge_capacity_Ah_mean",
    "discharge_capacity_Ah_std",
    "capacity_retention_pct_mean",
    "capacity_retention_pct_std",
    "energy_retention_pct_mean",
    "energy_efficiency_pct_mean",
    "coulombic_efficiency_pct_mean",
]].merge(sim_capacity_retention, on="cycle", how="left")

for sim_col in ["cw391_capacity_retention_pct", "cw511_capacity_retention_pct"]:
    if sim_col in benchmark_capacity:
        benchmark_capacity[f"{sim_col}_error_vs_exp_pctpt"] = (
            benchmark_capacity[sim_col] - benchmark_capacity["capacity_retention_pct_mean"]
        )

latest_exp = exp_cycles.sort_values("cycle").groupby(["group", "barcode"], as_index=False).tail(1)
final_summary = pd.concat(
    [
        pd.DataFrame({
            "scope": ["overall"],
            "group": ["all"],
            "barcode": ["all"],
            "cycle": [exp_overall_summary["cycle"].max()],
            "capacity_retention_pct": [
                exp_overall_summary.loc[
                    exp_overall_summary["cycle"].idxmax(),
                    "capacity_retention_pct_mean",
                ]
            ],
            "energy_retention_pct": [
                exp_overall_summary.loc[
                    exp_overall_summary["cycle"].idxmax(),
                    "energy_retention_pct_mean",
                ]
            ],
            "energy_efficiency_pct": [
                exp_overall_summary.loc[
                    exp_overall_summary["cycle"].idxmax(),
                    "energy_efficiency_pct_mean",
                ]
            ],
        }),
        latest_exp.rename(columns={
            "capacity_retention_pct": "capacity_retention_pct",
            "energy_retention_pct": "energy_retention_pct",
            "energy_efficiency_pct": "energy_efficiency_pct",
        })[[
            "group",
            "barcode",
            "cycle",
            "capacity_retention_pct",
            "energy_retention_pct",
            "energy_efficiency_pct",
        ]].assign(scope="cell")[[
            "scope",
            "group",
            "barcode",
            "cycle",
            "capacity_retention_pct",
            "energy_retention_pct",
            "energy_efficiency_pct",
        ]],
    ],
    ignore_index=True,
)

readme = pd.DataFrame({
    "item": [
        "source_workbook",
        "test_condition",
        "exported_file",
        "exp_cells",
        "exp_cycle_range",
        "sim_included",
    ],
    "value": [
        r"D:\Users\hez\Desktop\hithium\data_raw\314Ah\Exp\CW391+CW511-仿真支持数据 to 何争\软包-循环模版25℃-0.25P.xlsm",
        "25C 0.25P cycle benchmark",
        str(EXPORT_XLSX_PATH),
        int(exp_cycles["barcode"].nunique()),
        f"{int(exp_cycles['cycle'].min())}-{int(exp_cycles['cycle'].max())}",
        not sim_capacity_retention.empty,
    ],
})

with pd.ExcelWriter(EXPORT_XLSX_PATH, engine="openpyxl") as writer:
    readme.to_excel(writer, sheet_name="readme", index=False)
    exp_cycles.to_excel(writer, sheet_name="exp_cycles", index=False)
    exp_group_summary.to_excel(writer, sheet_name="exp_group_summary", index=False)
    exp_overall_summary.to_excel(writer, sheet_name="exp_overall_summary", index=False)
    sim_capacity_retention.to_excel(writer, sheet_name="sim_capacity_retention", index=False)
    benchmark_capacity.to_excel(writer, sheet_name="benchmark_capacity", index=False)
    final_summary.to_excel(writer, sheet_name="final_summary", index=False)

wb = load_workbook(EXPORT_XLSX_PATH)
header_fill = PatternFill("solid", fgColor="D9EAF7")
thin = Side(style="thin", color="D9D9D9")
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        width = min(max([len(v) for v in values] + [10]) + 2, 42)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
wb.save(EXPORT_XLSX_PATH)

print(f"Exported benchmark workbook: {EXPORT_XLSX_PATH}")
print("Sheets: readme, exp_cycles, exp_group_summary, exp_overall_summary, sim_capacity_retention, benchmark_capacity, final_summary")

