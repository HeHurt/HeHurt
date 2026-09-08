"""Run and report the 587 Ah, 0-50 C, 1-14 month unactivated storage matrix."""

from __future__ import annotations

import argparse
from datetime import datetime
import json
from pathlib import Path
import sys
import traceback
from typing import Iterable

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd


def locate_project_root(search_root: Path | None = None) -> Path:
    """Locate BatteryProject from a task Notebook, script, or repository path."""
    start = (search_root or Path.cwd()).resolve()
    for candidate in (start, *start.parents):
        if (candidate / "src" / "easy_imports.py").exists():
            return candidate
        nested = candidate / "BatteryProject"
        if (nested / "src" / "easy_imports.py").exists():
            return nested
    raise FileNotFoundError(f"Cannot locate BatteryProject from {start}")


PROJECT_ROOT = locate_project_root(Path(__file__).resolve())
WORKSPACE_ROOT = PROJECT_ROOT.parent
for import_root in (PROJECT_ROOT, WORKSPACE_ROOT):
    if str(import_root) not in sys.path:
        sys.path.insert(0, str(import_root))

from src.workflows.calendar_aging import CalendarAgingSpec  # noqa: E402
from src.workflows.calendar_aging_half_soc import (  # noqa: E402
    AGING_T_FACTOR,
    prepare_half_soc_baseline,
    run_half_soc_case_workflow,
)


FULL_TEMPERATURES_C = tuple(range(0, 51, 5))
FULL_MONTHS = tuple(range(1, 15))
SMOKE_TEMPERATURES_C = (0, 25, 50)
SMOKE_MONTHS = (1,)
CELL = "587calander"
NOMINAL_CAPACITY_AH = 587.0
NOMINAL_VOLTAGE_V = 3.2
DIAGNOSTIC_RATE_P = 0.5
DIAGNOSTIC_POWER_W = (
    NOMINAL_CAPACITY_AH * NOMINAL_VOLTAGE_V * DIAGNOSTIC_RATE_P
)
DAYS_PER_MONTH = 30.0


def _matrix_config(
    temperatures_c: Iterable[int],
    months: Iterable[int],
) -> dict:
    return {
        "cell": CELL,
        "aging_mode": "unactivated",
        "temperatures_c": list(temperatures_c),
        "months": list(months),
        "days_per_month": DAYS_PER_MONTH,
        "storage_soc_pct": 50,
        "test_temperature_c": 25,
        "relaxation_hours_at_25c": 1,
        "aging_t_factor": AGING_T_FACTOR,
        "diagnostic_rate_p": DIAGNOSTIC_RATE_P,
        "diagnostic_power_w": DIAGNOSTIC_POWER_W,
        "nominal_capacity_ah": NOMINAL_CAPACITY_AH,
        "nominal_voltage_v": NOMINAL_VOLTAGE_V,
        "charge_cutoff_v": 3.65,
        "discharge_cutoff_v": 2.5,
        "rest_period_hours": 24,
        "diagnostic_period_minutes": 0.5,
    }


def _write_json(path: Path, payload: dict) -> None:
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def _case_run_id(batch_id: str, temperature_c: int, month: int) -> str:
    return f"{batch_id}_T{temperature_c:02d}C_M{month:02d}"


def _case_metrics_path(
    project_root: Path,
    run_id: str,
) -> Path:
    return (
        project_root
        / "output"
        / "runs"
        / "calendar_aging_half_soc"
        / run_id
        / "artifacts"
        / "calendar_aging_half_soc_metrics.csv"
    )


def _resolve_case_run_id(
    project_root: Path,
    base_run_id: str,
) -> tuple[str, Path | None]:
    base_metrics = _case_metrics_path(project_root, base_run_id)
    if base_metrics.is_file():
        return base_run_id, base_metrics
    base_dir = base_metrics.parents[1]
    if not base_dir.exists():
        return base_run_id, None
    retry = 1
    while True:
        retry_id = f"{base_run_id}_retry{retry}"
        retry_metrics = _case_metrics_path(project_root, retry_id)
        if retry_metrics.is_file():
            return retry_id, retry_metrics
        if not retry_metrics.parents[1].exists():
            return retry_id, None
        retry += 1


def _case_spec(temperature_c: int, month: int) -> CalendarAgingSpec:
    return CalendarAgingSpec.from_mapping({
        "cell": CELL,
        "aging_mode": "unactivated",
        "run_mode": "study",
        "temperature_c": temperature_c,
        "diagnostic_rate_p": DIAGNOSTIC_RATE_P,
        "nominal_capacity_ah": NOMINAL_CAPACITY_AH,
        "nominal_voltage_v": NOMINAL_VOLTAGE_V,
        "unactivated_target_days": [month * DAYS_PER_MONTH],
        "rest_period_hours": 24,
        "diagnostic_period_minutes": 0.5,
        "charge_cutoff_v": 3.65,
        "discharge_cutoff_v": 2.5,
        "showprogress": False,
        "return_solutions": False,
    })


def _decorate_metrics(
    frame: pd.DataFrame,
    *,
    temperature_c: int,
    month: int,
    run_id: str,
) -> pd.DataFrame:
    result = frame.copy()
    result.insert(0, "temperature_c", temperature_c)
    result.insert(1, "storage_month", month)
    result["diagnostic_rate_p"] = DIAGNOSTIC_RATE_P
    result["diagnostic_power_w"] = DIAGNOSTIC_POWER_W
    result["storage_soc_pct"] = 50.0
    result["test_temperature_c"] = 25.0
    result["relaxation_hours_at_25c"] = 1.0
    result["aging_t_factor"] = AGING_T_FACTOR
    result["run_id"] = run_id
    return result


def _write_checkpoints(
    artifacts_dir: Path,
    metric_frames: list[pd.DataFrame],
    status_rows: list[dict],
) -> None:
    if metric_frames:
        metrics = pd.concat(metric_frames, ignore_index=True)
        metrics = metrics.sort_values(
            ["temperature_c", "storage_month"],
        ).reset_index(drop=True)
        metrics.to_csv(
            artifacts_dir / "calendar_aging_587Ah_50SOC_matrix.csv",
            index=False,
            encoding="utf-8-sig",
        )
    pd.DataFrame(status_rows).to_csv(
        artifacts_dir / "run_status.csv",
        index=False,
        encoding="utf-8-sig",
    )


def run_matrix(
    *,
    mode: str = "smoke",
    batch_id: str | None = None,
    project_root: Path | None = None,
) -> dict:
    """Run the requested matrix with one independently resumable run per case."""
    if mode not in {"smoke", "full"}:
        raise ValueError("mode must be 'smoke' or 'full'")
    project_root = (project_root or PROJECT_ROOT).resolve()
    temperatures = (
        SMOKE_TEMPERATURES_C if mode == "smoke" else FULL_TEMPERATURES_C
    )
    months = SMOKE_MONTHS if mode == "smoke" else FULL_MONTHS
    batch_id = batch_id or datetime.now().strftime(
        f"%Y%m%d_%H%M%S_587Ah_calendar_{mode}"
    )
    matrix_dir = (
        project_root
        / "output"
        / "runs"
        / "calendar_aging_matrix"
        / batch_id
    )
    artifacts_dir = matrix_dir / "artifacts"
    plots_dir = matrix_dir / "plots"
    artifacts_dir.mkdir(parents=True, exist_ok=True)
    plots_dir.mkdir(parents=True, exist_ok=True)
    config = _matrix_config(temperatures, months)
    config.update({"mode": mode, "batch_id": batch_id})
    _write_json(matrix_dir / "config.json", config)
    baseline = prepare_half_soc_baseline(
        _case_spec(25, 1),
        test_temperature_c=25,
    )
    pd.DataFrame([baseline.metrics()]).to_csv(
        artifacts_dir / "baseline_25C_C0.csv",
        index=False,
        encoding="utf-8-sig",
    )

    metric_frames: list[pd.DataFrame] = []
    status_rows: list[dict] = []
    total = len(temperatures) * len(months)
    completed = 0
    for temperature_c in temperatures:
        for month in months:
            completed += 1
            base_run_id = _case_run_id(batch_id, temperature_c, month)
            run_id, existing_metrics = _resolve_case_run_id(
                project_root,
                base_run_id,
            )
            print(
                f"[{completed}/{total}] {temperature_c}°C, month {month}, "
                f"run_id={run_id}",
                flush=True,
            )
            try:
                if existing_metrics is not None:
                    case_metrics = pd.read_csv(existing_metrics)
                    status = "reused"
                else:
                    result = run_half_soc_case_workflow(
                        _case_spec(temperature_c, month),
                        storage_days=month * DAYS_PER_MONTH,
                        baseline=baseline,
                        project_root=project_root,
                        run_id=run_id,
                        relaxation_hours=1,
                    )
                    case_metrics = result["metrics"]
                    status = "completed"
                metric_frames.append(
                    _decorate_metrics(
                        case_metrics,
                        temperature_c=temperature_c,
                        month=month,
                        run_id=run_id,
                    )
                )
                status_rows.append({
                    "temperature_c": temperature_c,
                    "storage_month": month,
                    "run_id": run_id,
                    "status": status,
                    "error": "",
                })
            except Exception as exc:
                error_path = artifacts_dir / (
                    f"error_T{temperature_c:02d}C_M{month:02d}.txt"
                )
                error_path.write_text(traceback.format_exc(), encoding="utf-8")
                status_rows.append({
                    "temperature_c": temperature_c,
                    "storage_month": month,
                    "run_id": run_id,
                    "status": "failed",
                    "error": f"{type(exc).__name__}: {exc}",
                })
                print(f"FAILED: {type(exc).__name__}: {exc}", flush=True)
            _write_checkpoints(artifacts_dir, metric_frames, status_rows)

    status = pd.DataFrame(status_rows)
    if not metric_frames:
        raise RuntimeError(
            f"all cases failed; see {artifacts_dir / 'run_status.csv'}"
        )
    metrics = (
        pd.concat(metric_frames, ignore_index=True)
        .sort_values(["temperature_c", "storage_month"])
        .reset_index(drop=True)
    )
    outputs = build_reports(
        metrics,
        status,
        artifacts_dir=artifacts_dir,
        plots_dir=plots_dir,
    )
    failures = status.loc[status["status"] == "failed"]
    if not failures.empty:
        raise RuntimeError(
            f"{len(failures)} cases failed; see {artifacts_dir / 'run_status.csv'}"
        )
    return {
        "batch_id": batch_id,
        "matrix_dir": matrix_dir,
        "artifacts_dir": artifacts_dir,
        "plots_dir": plots_dir,
        "metrics": metrics,
        "status": status,
        **outputs,
    }


def _pivot_percent(
    metrics: pd.DataFrame,
    value: str,
) -> pd.DataFrame:
    pivot = metrics.pivot(
        index="temperature_c",
        columns="storage_month",
        values=value,
    )
    pivot.index.name = "Temperature (°C)"
    pivot.columns = [f"Month {int(month)}" for month in pivot.columns]
    return pivot * 100


def _format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    thin = Side(style="thin", color="B7B7B7")
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for row in sheet.iter_rows():
            for cell in row:
                cell.border = Border(
                    left=thin,
                    right=thin,
                    top=thin,
                    bottom=thin,
                )
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )
        for column in sheet.columns:
            width = min(
                30,
                max(len(str(cell.value or "")) for cell in column) + 2,
            )
            sheet.column_dimensions[column[0].column_letter].width = width
        if (
            sheet.title
            in {
                "Retention_pct",
                "Recovery_pct",
                "Half_SOC_loss_pct",
                "Irreversible_loss_pct",
            }
            and sheet.max_row > 1
            and sheet.max_column > 1
        ):
            data_range = (
                f"B2:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
            )
            higher_is_better = sheet.title in {
                "Retention_pct",
                "Recovery_pct",
            }
            sheet.conditional_formatting.add(
                data_range,
                ColorScaleRule(
                    start_type="min",
                    start_color=(
                        "F8696B" if higher_is_better else "63BE7B"
                    ),
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color=(
                        "63BE7B" if higher_is_better else "F8696B"
                    ),
                ),
            )
    workbook.save(path)


def _plot_lines(metrics: pd.DataFrame, plots_dir: Path) -> Path:
    fig, axes = plt.subplots(1, 2, figsize=(13, 5), sharex=True)
    cmap = plt.get_cmap("coolwarm")
    norm = plt.Normalize(
        metrics["temperature_c"].min(),
        metrics["temperature_c"].max(),
    )
    for temperature_c, frame in metrics.groupby("temperature_c"):
        color = cmap(norm(temperature_c))
        label = f"{temperature_c:g}°C"
        axes[0].plot(
            frame["storage_month"],
            frame["retention_rate"] * 100,
            "o-",
            color=color,
            label=label,
        )
        axes[1].plot(
            frame["storage_month"],
            frame["recovery_rate"] * 100,
            "o-",
            color=color,
            label=label,
        )
    axes[0].set_title("Capacity Retention after Storage")
    axes[1].set_title("Recovered Capacity after Recharge")
    for axis in axes:
        axis.set_xlabel("Storage Month")
        axis.set_ylabel("Capacity (%)")
        axis.grid(True, alpha=0.3)
    axes[1].legend(ncol=2, fontsize=8)
    fig.tight_layout()
    path = plots_dir / "retention_recovery_vs_month.png"
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return path


def _plot_heatmap(
    metrics: pd.DataFrame,
    *,
    value: str,
    title: str,
    output_path: Path,
    cmap: str = "RdYlGn",
) -> Path:
    pivot = metrics.pivot(
        index="temperature_c",
        columns="storage_month",
        values=value,
    ) * 100
    fig, ax = plt.subplots(figsize=(11, 5.5))
    image = ax.imshow(pivot.values, aspect="auto", cmap=cmap)
    ax.set_xticks(np.arange(len(pivot.columns)), pivot.columns)
    ax.set_yticks(
        np.arange(len(pivot.index)),
        [f"{temperature:g}" for temperature in pivot.index],
    )
    ax.set_xlabel("Storage Month")
    ax.set_ylabel("Temperature (°C)")
    ax.set_title(title)
    fig.colorbar(image, ax=ax, label="Percent (%)")
    fig.tight_layout()
    fig.savefig(output_path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return output_path


def build_reports(
    metrics: pd.DataFrame,
    status: pd.DataFrame,
    *,
    artifacts_dir: Path,
    plots_dir: Path,
) -> dict:
    """Create the consolidated Excel workbook and summary plots."""
    csv_path = artifacts_dir / "calendar_aging_587Ah_50SOC_matrix.csv"
    metrics.to_csv(csv_path, index=False, encoding="utf-8-sig")
    final_month = int(metrics["storage_month"].max())
    final_month_summary = metrics.loc[
        metrics["storage_month"] == final_month
    ].copy()
    summary_path = artifacts_dir / f"summary_month_{final_month}.csv"
    final_month_summary.to_csv(
        summary_path,
        index=False,
        encoding="utf-8-sig",
    )
    workbook_path = (
        artifacts_dir
        / "587Ah_0to50C_50SOC_calendar_aging_14months.xlsx"
    )
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        metrics.to_excel(writer, sheet_name="All_results", index=False)
        _pivot_percent(metrics, "retention_rate").to_excel(
            writer,
            sheet_name="Retention_pct",
        )
        _pivot_percent(metrics, "recovery_rate").to_excel(
            writer,
            sheet_name="Recovery_pct",
        )
        _pivot_percent(
            metrics,
            "half_soc_capacity_loss_rate",
        ).to_excel(writer, sheet_name="Half_SOC_loss_pct")
        _pivot_percent(
            metrics,
            "irreversible_capacity_loss_rate",
        ).to_excel(writer, sheet_name="Irreversible_loss_pct")
        status.to_excel(writer, sheet_name="Run_status", index=False)
        final_month_summary.to_excel(
            writer,
            sheet_name=f"Month{final_month}_summary",
            index=False,
        )
    _format_workbook(workbook_path)
    line_plot = _plot_lines(metrics, plots_dir)
    retention_heatmap = _plot_heatmap(
        metrics,
        value="retention_rate",
        title="Capacity Retention: Temperature vs Storage Month",
        output_path=plots_dir / "retention_heatmap.png",
    )
    recovery_heatmap = _plot_heatmap(
        metrics,
        value="recovery_rate",
        title="Recovery Rate: Temperature vs Storage Month",
        output_path=plots_dir / "recovery_heatmap.png",
    )
    loss_heatmap = _plot_heatmap(
        metrics,
        value="irreversible_capacity_loss_rate",
        title="Irreversible Capacity Loss: Temperature vs Storage Month",
        output_path=plots_dir / "irreversible_loss_heatmap.png",
        cmap="RdYlGn_r",
    )
    return {
        "csv_path": csv_path,
        "summary_path": summary_path,
        "workbook_path": workbook_path,
        "plot_paths": [
            line_plot,
            retention_heatmap,
            recovery_heatmap,
            loss_heatmap,
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=("smoke", "full"), default="smoke")
    parser.add_argument("--run-id")
    args = parser.parse_args()
    result = run_matrix(mode=args.mode, batch_id=args.run_id)
    print(f"Matrix folder: {result['matrix_dir']}")
    print(f"Workbook: {result['workbook_path']}")


if __name__ == "__main__":
    main()
