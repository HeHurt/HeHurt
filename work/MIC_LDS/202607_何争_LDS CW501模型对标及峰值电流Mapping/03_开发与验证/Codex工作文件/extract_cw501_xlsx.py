import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def serialise(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


for file_name in sys.argv[1:]:
    path = Path(file_name)
    workbook = load_workbook(path, read_only=True, data_only=False)
    result = {"file": str(path), "sheets": []}
    for sheet in workbook.worksheets:
        rows = []
        for row_number, row in enumerate(sheet.iter_rows(), start=1):
            values = [serialise(cell.value) for cell in row]
            if any(value is not None for value in values):
                rows.append({"row": row_number, "values": values})
        result["sheets"].append(
            {
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "rows": rows,
            }
        )
    print(json.dumps(result, ensure_ascii=False))
