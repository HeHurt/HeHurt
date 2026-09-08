"""Validate prior CW363 discharge OCV and calibrate charge/discharge static OCV."""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import pybamm
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator


TASK_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = Path(__file__).resolve().parents[4]
PARAMS_ROOT = PROJECT_ROOT / "params"
RUNS_ROOT = PROJECT_ROOT / "BatteryProject" / "output" / "runs" / "cw363_soc_ocv"
SOURCE_XLSX = Path(
    r"E:\Downloads\Hithium_TI_∞Cell 1175Ah_不同温度、倍率SOC-OCV "
    r"SOC-OCV at different temp. & rate_20260707_V1.1.xlsx"
)
PRIOR_CURVE = (
    TASK_ROOT
    / "04_输出结果"
    / "20260827_140159_temperature_specific"
    / "temperature_specific_soc_ocv.csv"
)
TEMPERATURES = (25.0, 45.0)
POWER_W = 0.25 * 1175 * 3.2
REST_HOURS = 3
VAR_PTS = {"x_n": 5, "x_s": 5, "x_p": 5, "r_n": 20, "r_p": 20}

if str(PARAMS_ROOT) not in sys.path:
    sys.path.insert(0, str(PARAMS_ROOT))

from params1300CW363 import get_hithium_params  # noqa: E402


def load_source_measurements():
    workbook = load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    rows = []
    for sheet_name in ("10℃ 0.25P", "25℃ 0.25P", "25℃ 0.5P", "30℃ 0.25P", "45℃ 0.25P"):
        sheet = workbook[sheet_name]
        temperature_c = float(sheet_name.split("℃")[0])
        rate_p = float(sheet_name.split()[1].replace("P", ""))
        values = list(sheet.iter_rows(values_only=True))
        header_index = next(
            index for index, row in enumerate(values) if row[0] == "SOC/%" and row[2] == "SOC/%"
        )
        mappings = (
            (1, "discharge", 0, 1),
            (1, "charge", 2, 3),
            (2, "discharge", 5, 6),
            (2, "charge", 7, 8),
        )
        for row in values[header_index + 1 :]:
            for cell_id, direction, soc_col, voltage_col in mappings:
                soc = row[soc_col]
                voltage = row[voltage_col]
                if isinstance(soc, (int, float)) and isinstance(voltage, (int, float)):
                    rows.append(
                        {
                            "sheet": sheet_name,
                            "temperature_c": temperature_c,
                            "rate_p": rate_p,
                            "cell_id": cell_id,
                            "direction": direction,
                            "soc_pct": int(round(float(soc) * 100)),
                            "rested_voltage_v": float(voltage),
                        }
                    )
    workbook.close()
    raw = pd.DataFrame(rows)
    averaged = (
        raw.groupby(["temperature_c", "rate_p", "direction", "soc_pct"], as_index=False)
        .agg(
            exp_ocv_v=("rested_voltage_v", "mean"),
            cell_1_ocv_v=("rested_voltage_v", "first"),
            cell_2_ocv_v=("rested_voltage_v", "last"),
            cell_range_mv=("rested_voltage_v", lambda x: float((x.max() - x.min()) * 1000)),
        )
        .sort_values(["temperature_c", "rate_p", "direction", "soc_pct"])
    )
    return raw, averaged


def scalar_parameter(parameters, key, sto):
    value = parameters[key]
    symbol = value(pybamm.Scalar(float(sto))) if callable(value) else pybamm.Scalar(value)
    return float(np.asarray(parameters.process_symbol(symbol).evaluate()).reshape(-1)[0])


def fresh_parameters(temperature_k, charge_ocp=None, discharge_ocp=None):
    parameters = pybamm.ParameterValues("OKane2022")
    parameters.update(
        get_hithium_params(t_factor=1, temperature=temperature_k),
        check_already_exists=False,
    )
    parameters.update(
        {
            "Ambient temperature [K]": temperature_k,
            "Initial temperature [K]": temperature_k,
        },
        check_already_exists=False,
    )
    updates = {}
    if charge_ocp is not None:
        updates["Positive electrode delithiation OCP [V]"] = charge_ocp
    if discharge_ocp is not None:
        updates["Positive electrode OCP [V]"] = discharge_ocp
        updates["Positive electrode lithiation OCP [V]"] = discharge_ocp
    if updates:
        parameters.update(updates)
    return parameters


def last_value(solution, name):
    return float(np.asarray(solution[name].entries).reshape(-1)[-1])


def state_from_rest(solution, parameters):
    cn = last_value(solution, "X-averaged negative particle concentration [mol.m-3]")
    cp = last_value(solution, "X-averaged positive particle concentration [mol.m-3]")
    return {
        "xn": cn / float(parameters["Maximum concentration in negative electrode [mol.m-3]"]),
        "xp": cp / float(parameters["Maximum concentration in positive electrode [mol.m-3]"]),
        "terminal_voltage_v": last_value(solution, "Terminal voltage [V]"),
        "midpoint_ocv_v": last_value(solution, "Battery open-circuit voltage [V]"),
    }


def run_endpoint_protocol(temperature_c, charge_ocp=None, discharge_ocp=None):
    temperature_k = temperature_c + 273.15
    parameters = fresh_parameters(temperature_k, charge_ocp, discharge_ocp)
    model = pybamm.lithium_ion.DFN(
        {
            "calculate discharge energy": "true",
            "contact resistance": "true",
            "open-circuit potential": ("current sigmoid", "current sigmoid"),
        }
    )
    experiment = pybamm.Experiment(
        [
            f"Charge at {POWER_W:g} W until 3.65 V",
            f"Rest for {REST_HOURS} hours",
            f"Discharge at {POWER_W:g} W until 2.50 V",
            f"Rest for {REST_HOURS} hours",
        ],
        period="60 seconds",
    )
    simulation = pybamm.Simulation(
        model,
        parameter_values=parameters,
        experiment=experiment,
        var_pts=VAR_PTS,
        solver=pybamm.IDAKLUSolver(rtol=1e-6, atol=1e-6),
    )
    solution = simulation.solve(initial_soc=0.5)
    steps = [step for cycle in solution.cycles for step in cycle.steps]
    if len(steps) != 4:
        raise RuntimeError(f"Expected four protocol steps, got {len(steps)}")
    charge_step, rest_100, discharge_step, rest_0 = steps
    state_100 = state_from_rest(rest_100, parameters)
    state_0 = state_from_rest(rest_0, parameters)
    discharge_t = np.asarray(discharge_step["Time [s]"].entries).reshape(-1)
    discharge_i = np.asarray(discharge_step["Current [A]"].entries).reshape(-1)
    endpoint = {
        "temperature_c": temperature_c,
        "xn_0": state_0["xn"],
        "xp_0": state_0["xp"],
        "rested_terminal_0_v": state_0["terminal_voltage_v"],
        "midpoint_ocv_0_v": state_0["midpoint_ocv_v"],
        "xn_100": state_100["xn"],
        "xp_100": state_100["xp"],
        "rested_terminal_100_v": state_100["terminal_voltage_v"],
        "midpoint_ocv_100_v": state_100["midpoint_ocv_v"],
        "discharge_capacity_ah": float(np.trapz(discharge_i, discharge_t) / 3600),
    }
    timeseries = pd.DataFrame(
        {
            "temperature_c": temperature_c,
            "time_h": np.asarray(solution["Time [s]"].entries).reshape(-1) / 3600,
            "terminal_voltage_v": np.asarray(solution["Terminal voltage [V]"].entries).reshape(-1),
        }
    )
    return parameters, endpoint, timeseries


def branch_ocv(parameters, endpoints, soc, temperature_k, direction):
    xn = endpoints["xn_0"] + soc * (endpoints["xn_100"] - endpoints["xn_0"])
    xp = endpoints["xp_0"] + soc * (endpoints["xp_100"] - endpoints["xp_0"])
    if direction == "charge":
        up_key = "Positive electrode delithiation OCP [V]"
        un_key = "Negative electrode lithiation OCP [V]"
    else:
        up_key = "Positive electrode lithiation OCP [V]"
        un_key = "Negative electrode delithiation OCP [V]"
    reference_k = float(parameters["Reference temperature [K]"])
    up = scalar_parameter(parameters, up_key, xp)
    un = scalar_parameter(parameters, un_key, xn)
    dup_dt = scalar_parameter(parameters, "Positive electrode OCP entropic change [V.K-1]", xp)
    dun_dt = scalar_parameter(parameters, "Negative electrode OCP entropic change [V.K-1]", xn)
    ocv = up - un + (temperature_k - reference_k) * (dup_dt - dun_dt)
    return ocv, xn, xp


def build_positive_ocp(temperature_c, endpoints, target, direction):
    temperature_k = temperature_c + 273.15
    base_parameters = fresh_parameters(temperature_k)
    positive_key = (
        "Positive electrode delithiation OCP [V]"
        if direction == "charge"
        else "Positive electrode lithiation OCP [V]"
    )
    base_positive = base_parameters[positive_key]
    soc_grid = np.linspace(0, 1, 1001)
    sto, correction = [], []
    for soc in soc_grid:
        base_ocv, _, xp = branch_ocv(base_parameters, endpoints, soc, temperature_k, direction)
        sto.append(xp)
        correction.append(float(target(soc)) - base_ocv)
    order = np.argsort(sto)
    sto_knots = np.asarray(sto)[order]
    correction_v = np.asarray(correction)[order]
    sto_knots = np.concatenate(([0.0], sto_knots, [1.0]))
    correction_v = np.concatenate(([correction_v[0]], correction_v, [correction_v[-1]]))

    def calibrated_ocp(stoichiometry):
        delta = pybamm.Interpolant(
            sto_knots,
            correction_v,
            stoichiometry,
            name=f"CW363 {temperature_c:.0f}C {direction} OCP correction",
            interpolator="linear",
        )
        return base_positive(stoichiometry) + delta

    return calibrated_ocp


def endpoint_delta(old, new):
    return max(
        abs(float(old[key]) - float(new[key]))
        for key in ("xn_0", "xp_0", "xn_100", "xp_100")
    )


def calibrate_temperature(temperature_c, measured_average, max_iterations=12, tolerance=1e-4):
    targets = {}
    for direction in ("charge", "discharge"):
        subset = measured_average[
            (measured_average["temperature_c"] == temperature_c)
            & (measured_average["rate_p"] == 0.25)
            & (measured_average["direction"] == direction)
        ].sort_values("soc_pct")
        targets[direction] = PchipInterpolator(
            subset["soc_pct"].to_numpy(dtype=float) / 100,
            subset["exp_ocv_v"].to_numpy(dtype=float),
        )
    _, endpoints, _ = run_endpoint_protocol(temperature_c)
    history = []
    parameters = None
    timeseries = None
    charge_ocp = None
    discharge_ocp = None
    for iteration in range(1, max_iterations + 1):
        charge_ocp = build_positive_ocp(temperature_c, endpoints, targets["charge"], "charge")
        discharge_ocp = build_positive_ocp(
            temperature_c, endpoints, targets["discharge"], "discharge"
        )
        parameters, new_endpoints, timeseries = run_endpoint_protocol(
            temperature_c, charge_ocp, discharge_ocp
        )
        delta = endpoint_delta(endpoints, new_endpoints)
        history.append({"iteration": iteration, "endpoint_max_abs_delta": delta})
        endpoints = new_endpoints
        if delta < tolerance:
            break
    # Rebuild the static OCP tables at the final protocol-defined endpoints.
    # The last dynamic run remains the endpoint audit; this final rebuild makes
    # the exported static SOC-OCV curves correspond exactly to those endpoints.
    charge_ocp = build_positive_ocp(temperature_c, endpoints, targets["charge"], "charge")
    discharge_ocp = build_positive_ocp(
        temperature_c, endpoints, targets["discharge"], "discharge"
    )
    parameters = fresh_parameters(temperature_c + 273.15, charge_ocp, discharge_ocp)
    return parameters, endpoints, timeseries, charge_ocp, discharge_ocp, targets, history


def build_output_curves(calibrations):
    rows = []
    ocp_rows = []
    for temperature_c, calibration in calibrations.items():
        parameters = calibration["parameters"]
        endpoints = calibration["endpoints"]
        targets = calibration["targets"]
        temperature_k = temperature_c + 273.15
        base_parameters = fresh_parameters(temperature_k)
        for direction in ("charge", "discharge"):
            for soc_pct in range(101):
                soc = soc_pct / 100
                calibrated, xn, xp = branch_ocv(
                    parameters, endpoints, soc, temperature_k, direction
                )
                baseline, _, _ = branch_ocv(
                    base_parameters, endpoints, soc, temperature_k, direction
                )
                rows.append(
                    {
                        "temperature_c": temperature_c,
                        "direction": direction,
                        "soc_pct": soc_pct,
                        "xn": xn,
                        "xp": xp,
                        "baseline_ocv_v": baseline,
                        "calibrated_ocv_v": calibrated,
                        "target_ocv_v": float(targets[direction](soc)),
                        "ocv_correction_mv": (calibrated - baseline) * 1000,
                    }
                )
            positive_key = (
                "Positive electrode delithiation OCP [V]"
                if direction == "charge"
                else "Positive electrode lithiation OCP [V]"
            )
            for xp in np.linspace(0, 1, 1001):
                ocp_rows.append(
                    {
                        "temperature_c": temperature_c,
                        "direction": direction,
                        "positive_stoichiometry": xp,
                        "calibrated_positive_ocp_v": scalar_parameter(parameters, positive_key, xp),
                    }
                )
    return pd.DataFrame(rows), pd.DataFrame(ocp_rows)


def error_metrics(data, error_column, group_columns):
    return (
        data.groupby(group_columns, as_index=False)
        .agg(
            points=(error_column, "count"),
            rmse_mv=(error_column, lambda x: float(np.sqrt(np.mean(np.asarray(x) ** 2)))),
            mean_error_mv=(error_column, "mean"),
            max_abs_error_mv=(error_column, lambda x: float(np.max(np.abs(x)))),
        )
        .sort_values(group_columns)
    )


def plot_prior_validation(prior_comparison, output_path):
    plt.style.use("science")
    plt.rcParams["font.family"] = ["Calibri", "Microsoft YaHei"]
    fig, axes = plt.subplots(2, 2, figsize=(11.2, 7.6), sharex="col")
    colors = {25.0: "#005BAC", 45.0: "#D1495B"}
    for column, temperature_c in enumerate(TEMPERATURES):
        subset = prior_comparison[prior_comparison["temperature_c"] == temperature_c]
        axes[0, column].plot(
            subset["soc_pct"], subset["prior_discharge_ocv_v"], color=colors[temperature_c], label="Prior Sim"
        )
        axes[0, column].scatter(
            subset["soc_pct"], subset["exp_ocv_v"], facecolors="none", edgecolors=colors[temperature_c], label="Full Exp"
        )
        axes[0, column].set(title=f"{temperature_c:.0f}°C Discharge Static OCV", ylabel="OCV (V)")
        axes[0, column].grid(alpha=0.25)
        axes[0, column].legend()
        axes[1, column].axhline(0, color="#555555", linewidth=0.8)
        axes[1, column].plot(subset["soc_pct"], subset["prior_error_mv"], color=colors[temperature_c], marker="o", markersize=3)
        axes[1, column].set(xlabel="SOC (%)", ylabel="Sim - Exp (mV)")
        axes[1, column].grid(alpha=0.25)
    fig.tight_layout()
    fig.savefig(output_path, dpi=240)
    plt.close(fig)


def plot_charge_discharge(comparison, output_path):
    plt.style.use("science")
    plt.rcParams["font.family"] = ["Calibri", "Microsoft YaHei"]
    fig, axes = plt.subplots(1, 2, figsize=(12.4, 5.1), sharey=True)
    colors = {"charge": "#1B75BC", "discharge": "#D1495B"}
    for ax, temperature_c in zip(axes, TEMPERATURES):
        subset = comparison[comparison["temperature_c"] == temperature_c]
        for direction in ("charge", "discharge"):
            direction_data = subset[subset["direction"] == direction].sort_values("soc_pct")
            ax.plot(
                direction_data["soc_pct"],
                direction_data["calibrated_ocv_v"],
                color=colors[direction],
                label=f"{direction.title()} Sim",
            )
            ax.scatter(
                direction_data["soc_pct"],
                direction_data["exp_ocv_v"],
                facecolors="none",
                edgecolors=colors[direction],
                label=f"{direction.title()} Exp",
            )
        ax.set(title=f"{temperature_c:.0f}°C 0.25P Static SOC-OCV", xlabel="SOC (%)", ylabel="OCV (V)")
        ax.grid(alpha=0.25)
        ax.legend(fontsize=8)
    fig.subplots_adjust(left=0.08, right=0.98, bottom=0.13, top=0.88, wspace=0.18)
    fig.savefig(output_path, dpi=240)
    plt.close(fig)


def plot_all_measured(averaged, output_path):
    data = averaged[averaged["rate_p"] == 0.25]
    plt.style.use("science")
    plt.rcParams["font.family"] = ["Calibri", "Microsoft YaHei"]
    fig, axes = plt.subplots(1, 2, figsize=(12.4, 5.1), sharey=True)
    colors = {10.0: "#6A3D9A", 25.0: "#005BAC", 30.0: "#2CA25F", 45.0: "#D1495B"}
    for ax, direction in zip(axes, ("charge", "discharge")):
        for temperature_c in (10.0, 25.0, 30.0, 45.0):
            subset = data[(data["temperature_c"] == temperature_c) & (data["direction"] == direction)]
            ax.plot(subset["soc_pct"], subset["exp_ocv_v"], marker="o", markersize=3, color=colors[temperature_c], label=f"{temperature_c:.0f}°C")
        ax.set(title=f"Measured {direction.title()} Static OCV", xlabel="SOC (%)", ylabel="OCV (V)")
        ax.grid(alpha=0.25)
        ax.legend(fontsize=8)
    fig.subplots_adjust(left=0.08, right=0.98, bottom=0.13, top=0.88, wspace=0.18)
    fig.savefig(output_path, dpi=240)
    plt.close(fig)


def main():
    raw, averaged = load_source_measurements()
    prior = pd.read_csv(PRIOR_CURVE)
    prior_discharge = prior[prior["temperature_c"].isin(TEMPERATURES)][
        ["temperature_c", "soc_pct", "temperature_specific_ocv_v"]
    ].rename(columns={"temperature_specific_ocv_v": "prior_discharge_ocv_v"})
    prior_exp = averaged[
        (averaged["temperature_c"].isin(TEMPERATURES))
        & (averaged["rate_p"] == 0.25)
        & (averaged["direction"] == "discharge")
    ]
    prior_comparison = prior_exp.merge(prior_discharge, on=["temperature_c", "soc_pct"])
    prior_comparison["prior_error_mv"] = (
        prior_comparison["prior_discharge_ocv_v"] - prior_comparison["exp_ocv_v"]
    ) * 1000
    prior_metrics = error_metrics(prior_comparison, "prior_error_mv", ["temperature_c"])

    calibrations = {}
    for temperature_c in TEMPERATURES:
        parameters, endpoints, timeseries, charge_ocp, discharge_ocp, targets, history = (
            calibrate_temperature(temperature_c, averaged)
        )
        calibrations[temperature_c] = {
            "parameters": parameters,
            "endpoints": endpoints,
            "timeseries": timeseries,
            "charge_ocp": charge_ocp,
            "discharge_ocp": discharge_ocp,
            "targets": targets,
            "history": history,
        }

    curves, ocp_tables = build_output_curves(calibrations)
    measured_primary = averaged[
        (averaged["temperature_c"].isin(TEMPERATURES)) & (averaged["rate_p"] == 0.25)
    ]
    comparison = measured_primary.merge(
        curves,
        on=["temperature_c", "direction", "soc_pct"],
        how="left",
    )
    comparison["error_mv"] = (comparison["calibrated_ocv_v"] - comparison["exp_ocv_v"]) * 1000
    calibrated_metrics = error_metrics(
        comparison, "error_mv", ["temperature_c", "direction"]
    )
    hysteresis = comparison.pivot_table(
        index=["temperature_c", "soc_pct"], columns="direction", values="exp_ocv_v"
    ).reset_index()
    hysteresis["charge_minus_discharge_mv"] = (
        hysteresis["charge"] - hysteresis["discharge"]
    ) * 1000

    run_dir = RUNS_ROOT / datetime.now().strftime("%Y%m%d_%H%M%S_full_charge_discharge")
    run_dir.mkdir(parents=True, exist_ok=False)
    raw.to_csv(run_dir / "source_measurements_long.csv", index=False, encoding="utf-8-sig")
    averaged.to_csv(run_dir / "source_measurements_average.csv", index=False, encoding="utf-8-sig")
    prior_comparison.to_csv(run_dir / "prior_discharge_validation.csv", index=False, encoding="utf-8-sig")
    prior_metrics.to_csv(run_dir / "prior_discharge_metrics.csv", index=False, encoding="utf-8-sig")
    curves.to_csv(run_dir / "charge_discharge_soc_ocv_1pct.csv", index=False, encoding="utf-8-sig")
    comparison.to_csv(run_dir / "charge_discharge_measured_comparison.csv", index=False, encoding="utf-8-sig")
    calibrated_metrics.to_csv(run_dir / "charge_discharge_metrics.csv", index=False, encoding="utf-8-sig")
    hysteresis.to_csv(run_dir / "measured_hysteresis_gap.csv", index=False, encoding="utf-8-sig")
    endpoint_rows = []
    timeseries_rows = []
    iteration_history = {}
    for temperature_c, calibration in calibrations.items():
        endpoint_rows.append(calibration["endpoints"])
        timeseries_rows.append(calibration["timeseries"])
        iteration_history[f"{temperature_c:.0f}C"] = calibration["history"]
        for direction in ("charge", "discharge"):
            table = ocp_tables[
                (ocp_tables["temperature_c"] == temperature_c)
                & (ocp_tables["direction"] == direction)
            ]
            table[["positive_stoichiometry", "calibrated_positive_ocp_v"]].to_csv(
                run_dir / f"LFP_CW363_{direction}_calibrated_{temperature_c:.0f}C.csv",
                index=False,
                header=False,
                encoding="utf-8",
            )
    endpoints = pd.DataFrame(endpoint_rows)
    endpoints.to_csv(run_dir / "protocol_endpoint_stoichiometry.csv", index=False, encoding="utf-8-sig")
    pd.concat(timeseries_rows, ignore_index=True).to_csv(
        run_dir / "endpoint_protocol_timeseries.csv", index=False, encoding="utf-8-sig"
    )
    plot_prior_validation(prior_comparison, run_dir / "prior_discharge_full_soc_validation.png")
    plot_charge_discharge(comparison, run_dir / "charge_discharge_ocv_comparison.png")
    plot_all_measured(averaged, run_dir / "all_temperature_measured_ocv.png")

    rate_check = averaged[
        (averaged["temperature_c"] == 25.0) & (averaged["rate_p"].isin([0.25, 0.5]))
    ].pivot_table(index=["direction", "soc_pct"], columns="rate_p", values="exp_ocv_v").reset_index()
    rate_check["0p5_minus_0p25_mv"] = (rate_check[0.5] - rate_check[0.25]) * 1000
    rate_check.to_csv(run_dir / "measured_rate_sensitivity_25C.csv", index=False, encoding="utf-8-sig")

    manifest = {
        "source": str(SOURCE_XLSX),
        "source_protocol": "0.25P steps with 3 h rest at each SOC point",
        "replicates": 2,
        "model_temperatures_c": list(TEMPERATURES),
        "power_w": POWER_W,
        "endpoint_tolerance": 1e-4,
        "static_ocp_rebuilt_at_final_protocol_endpoints": True,
        "prior_discharge_metrics": prior_metrics.to_dict(orient="records"),
        "calibrated_metrics": calibrated_metrics.to_dict(orient="records"),
        "endpoint_iteration_history": iteration_history,
        "pybamm_version": pybamm.__version__,
    }
    (run_dir / "run_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(run_dir)
    print("\nPrior discharge validation:\n", prior_metrics.to_string(index=False))
    print("\nCharge/discharge calibrated metrics:\n", calibrated_metrics.to_string(index=False))
    print("\nEndpoints:\n", endpoints.to_string(index=False))
    print("\n25C 0.5P - 0.25P measured range (mV):", float(rate_check["0p5_minus_0p25_mv"].min()), float(rate_check["0p5_minus_0p25_mv"].max()))


if __name__ == "__main__":
    main()
