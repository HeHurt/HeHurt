import json
from pathlib import Path

import matplotlib.pyplot as plt
import matplotlib as mpl
from matplotlib.colors import LogNorm
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


TASK = Path(r"C:\HithiumSSD\hithium\work\AI_virtual_cell\202609_何争_扣电与软包力电耦合V2")
DATA = TASK / "99_内部文件" / "中间数据"
SUMMARY = TASK / "99_内部文件" / "JSON" / "扣电V2_0p5C_完整循环后处理汇总.json"
OUT = TASK / "04_输出结果"
FIG = OUT / "图表"
FIG.mkdir(parents=True, exist_ok=True)

mpl.rcParams.update({
    "font.sans-serif": ["Microsoft YaHei", "SimHei", "Arial Unicode MS"],
    "axes.unicode_minus": False,
    "figure.dpi": 150,
    "savefig.dpi": 220,
    "axes.grid": True,
    "grid.alpha": 0.25,
    "axes.spines.top": False,
    "axes.spines.right": False,
})

payload = json.loads(SUMMARY.read_text(encoding="utf-8"))
metrics = payload["metrics"]
df = pd.read_csv(DATA / "扣电V2_0p5C_完整循环时序.csv")
t_h = df["time_s"] / 3600.0


def save(fig, name):
    fig.tight_layout()
    fig.savefig(FIG / name, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def phase_span(ax):
    ce = metrics["charge_end_time_s"] / 3600
    ds = metrics["discharge_start_time_s"] / 3600
    de = metrics["discharge_end_time_s"] / 3600
    ax.axvspan(0, ce, color="#4C78A8", alpha=0.07)
    ax.axvspan(ce, ds, color="#999999", alpha=0.08)
    ax.axvspan(ds, de, color="#E45756", alpha=0.06)
    ax.axvspan(de, t_h.iloc[-1], color="#999999", alpha=0.08)


# 01 完整循环
fig, ax = plt.subplots(figsize=(11, 5.6))
phase_span(ax)
ax.plot(t_h, df["terminal_voltage_V"], color="#1455C0", lw=2.0, label="负载端电压")
ax.plot(t_h, df["cell_voltage_V"], color="#7A9ED6", lw=1.0, ls="--", label="电化学端电压")
ax.axhline(3.65, color="crimson", ls=":", lw=1.3, label="3.65 V 充电截止")
ax.axhline(2.50, color="darkorange", ls=":", lw=1.3, label="2.50 V 放电截止")
ax.set(xlabel="时间 (h)", ylabel="电压 (V)", title="扣电 V2：100 N、0.5C 完整充放电循环")
ax2 = ax.twinx()
ax2.plot(t_h, df["current_Crate"], color="#333333", lw=1.1, alpha=0.75, label="倍率")
ax2.set_ylabel("电流倍率 (C)")
ax2.grid(False)
lines = ax.get_lines() + ax2.get_lines()
ax.legend(lines, [x.get_label() for x in lines], loc="lower center", ncol=3, frameon=True)
save(fig, "01_完整循环电压电流与阶段.png")


# 02 SOC 与嵌锂状态
fig, axes = plt.subplots(1, 2, figsize=(12, 4.8))
axes[0].plot(t_h, 100 * df["actual_SOC"], color="#1455C0", lw=2, label="库仑积分 SOC")
axes[0].plot(t_h, 100 * df["theta_negative"], color="#222222", lw=1.5, label="石墨平均嵌锂比")
axes[0].plot(t_h, 100 * (1 - df["theta_positive"]), color="#D84A3A", lw=1.5, label="LFP 脱锂度 1−θp")
axes[0].set(xlabel="时间 (h)", ylabel="状态 (%)", title="循环状态演化")
axes[0].legend(frameon=True)
charge = df[df["current_Crate"] > 0.1]
discharge = df[df["current_Crate"] < -0.1]
axes[1].plot(100 * charge["actual_SOC"], charge["terminal_voltage_V"], color="#1455C0", lw=2, label="充电")
axes[1].plot(100 * discharge["actual_SOC"], discharge["terminal_voltage_V"], color="#D84A3A", lw=2, label="放电")
axes[1].axhline(3.65, color="crimson", ls=":", lw=1)
axes[1].axhline(2.50, color="darkorange", ls=":", lw=1)
axes[1].set(xlabel="库仑积分 SOC (%)", ylabel="端电压 (V)", title="电压–SOC 滞回")
axes[1].legend(frameon=True)
save(fig, "02_SOC与正负极嵌锂状态.png")


# 03 压力、孔隙率与呼吸位移
fig, axes = plt.subplots(2, 2, figsize=(12, 8))
for col, label in [("pressure_negative_MPa", "负极"), ("pressure_separator_MPa", "隔膜"), ("pressure_positive_MPa", "正极")]:
    axes[0, 0].plot(t_h, df[col], lw=1.7, label=label)
axes[0, 0].set(xlabel="时间 (h)", ylabel="平均压缩压力 (MPa)", title="局部压力反馈")
axes[0, 0].legend(frameon=True)
for col, label in [("porosity_negative", "负极"), ("porosity_separator", "隔膜"), ("porosity_positive", "正极")]:
    axes[0, 1].plot(t_h, 100 * df[col], lw=1.7, label=label)
axes[0, 1].set(xlabel="时间 (h)", ylabel="孔隙率 (%)", title="压力反馈后的孔隙率")
axes[0, 1].legend(frameon=True)
axes[1, 0].plot(t_h, 100 * df["breathing_strain_negative"], lw=1.8, color="#222222", label="石墨")
axes[1, 0].plot(t_h, 100 * df["breathing_strain_positive"], lw=1.8, color="#D84A3A", label="LFP")
axes[1, 0].set(xlabel="时间 (h)", ylabel="本征呼吸应变 (%)", title="非线性 SOC–呼吸关系")
axes[1, 0].legend(frameon=True)
disp_rel = df["loaded_boundary_displacement_um"] - df["loaded_boundary_displacement_um"].iloc[0]
axes[1, 1].plot(t_h, disp_rel, lw=2, color="#7A3E9D")
axes[1, 1].set(xlabel="时间 (h)", ylabel="相对位移 (μm)", title="加载边界相对位移响应")
save(fig, "03_压力孔隙率呼吸与位移.png")


# 04 极化分解表征
fig, axes = plt.subplots(1, 2, figsize=(12, 4.8))
axes[0].plot(t_h, df["electrolyte_c_negative_mol_m3"], lw=1.7, label="负极")
axes[0].plot(t_h, df["electrolyte_c_separator_mol_m3"], lw=1.7, label="隔膜")
axes[0].plot(t_h, df["electrolyte_c_positive_mol_m3"], lw=1.7, label="正极")
axes[0].fill_between(t_h, df["electrolyte_c_min_mol_m3"], df["electrolyte_c_max_mol_m3"], color="#4C78A8", alpha=0.14, label="全域 min–max")
axes[0].set(xlabel="时间 (h)", ylabel="电解液浓度 (mol/m³)", title="浓差极化表征")
axes[0].legend(frameon=True)
axes[1].plot(t_h, df["contact_voltage_drop_mV"], lw=1.8, color="#D84A3A", label="接触压降")
axes[1].set(xlabel="时间 (h)", ylabel="接触压降 (mV)", title="接触电阻贡献")
axr = axes[1].twinx()
axr.plot(t_h, df["electrolyte_span_mol_m3"], lw=1.4, color="#1455C0", alpha=0.8, label="浓度跨度")
axr.set_ylabel("电解液浓度跨度 (mol/m³)")
axr.grid(False)
lines = axes[1].get_lines() + axr.get_lines()
axes[1].legend(lines, [x.get_label() for x in lines], loc="upper center", frameon=True)
save(fig, "04_浓差极化与接触压降.png")


# 05 集流构件电流密度空间分布
stages = [(phase, soc) for phase in ("充电", "放电") for soc in (10, 30, 50, 70, 90)]
all_current = []
current_data = {}
for phase, soc in stages:
    item = pd.read_csv(DATA / f"集流构件电流密度_{phase}_SOC{soc:02d}.csv")
    current_data[(phase, soc)] = item
    all_current.append(item.loc[item["current_density_A_m2"] > 0, "current_density_A_m2"].to_numpy())
positive_current = np.concatenate(all_current)
vmin = max(np.percentile(positive_current, 2), 1e-4)
vmax = np.percentile(positive_current, 99.7)
fig, axes = plt.subplots(2, 5, figsize=(17, 6.5), sharex=True, sharey=True)
for row, phase in enumerate(("充电", "放电")):
    for col, soc in enumerate((10, 30, 50, 70, 90)):
        item = current_data[(phase, soc)].iloc[::3]
        values = np.clip(item["current_density_A_m2"], vmin, None)
        sc = axes[row, col].scatter(item["r_um"] / 1000, item["z_um"] / 1000, c=values, s=2.5, cmap="turbo", norm=LogNorm(vmin=vmin, vmax=vmax), rasterized=True)
        axes[row, col].set_title(f"{phase} SOC {soc}%")
        if row == 1:
            axes[row, col].set_xlabel("径向 r (mm)")
        if col == 0:
            axes[row, col].set_ylabel("轴向 z (mm)")
cax = fig.add_axes([0.925, 0.14, 0.014, 0.70])
fig.colorbar(sc, cax=cax, label="固相电流密度 |Is| (A/m²)")
fig.suptitle("集流体与壳体电流密度分布（统一对数色标；轴向显示比例已放大）", fontsize=14)
fig.subplots_adjust(left=0.06, right=0.90, bottom=0.10, top=0.88, wspace=0.18, hspace=0.28)
fig.savefig(FIG / "05_集流构件电流密度_SOC阶段.png", bbox_inches="tight", facecolor="white")
plt.close(fig)


# 06 界面嵌锂状态径向分布
fig, axes = plt.subplots(2, 2, figsize=(12, 8), sharex=True)
interface_groups = [(["负极-隔膜", "负极-铜箔"], "负极"), (["正极-隔膜", "正极-铝箔"], "正极")]
colors = plt.cm.viridis(np.linspace(0.08, 0.92, 5))
for row, phase in enumerate(("充电", "放电")):
    for col, (interfaces, electrode) in enumerate(interface_groups):
        ax = axes[row, col]
        for color, soc in zip(colors, (10, 30, 50, 70, 90)):
            item = pd.read_csv(DATA / f"界面嵌锂状态_{phase}_SOC{soc:02d}.csv")
            for interface in interfaces:
                part = item[item["interface"] == interface].groupby("r_um", as_index=False)["electrode_stoichiometry"].mean()
                style = "-" if "隔膜" in interface else "--"
                ax.plot(part["r_um"] / 1000, part["electrode_stoichiometry"], color=color, ls=style, lw=1.3, label=f"SOC {soc}%·{interface.split('-')[1]}")
        ax.set_title(f"{phase}：{electrode}两侧界面")
        ax.set_ylabel("局部嵌锂比例 θ")
        if row == 1:
            ax.set_xlabel("径向 r (mm)")
        ax.legend(fontsize=7, ncol=2, frameon=True)
save(fig, "06_正负极界面嵌锂状态_SOC阶段.png")


# 07 局部压力和孔隙率径向分布
chosen = [("充电", 10), ("充电", 50), ("充电", 90), ("放电", 50), ("放电", 10)]
fig, axes = plt.subplots(2, 3, figsize=(15, 8), sharex=True)
for col, region in enumerate(("负极", "隔膜", "正极")):
    for (phase, soc), color in zip(chosen, plt.cm.plasma(np.linspace(0.08, 0.9, len(chosen)))):
        item = pd.read_csv(DATA / f"局部压力孔隙率_{phase}_SOC{soc:02d}.csv")
        part = item[item["region"] == region].copy()
        part["r_round"] = part["r_um"].round(3)
        radial = part.groupby("r_round", as_index=False).agg(pressure_Pa=("pressure_Pa", "mean"), porosity=("porosity", "mean"))
        label = f"{phase}{soc}%"
        axes[0, col].plot(radial["r_round"] / 1000, radial["pressure_Pa"] / 1e6, color=color, lw=1.4, label=label)
        axes[1, col].plot(radial["r_round"] / 1000, 100 * radial["porosity"], color=color, lw=1.4, label=label)
    axes[0, col].set_title(region)
    axes[1, col].set_xlabel("径向 r (mm)")
    axes[0, col].legend(fontsize=8, frameon=True)
axes[0, 0].set_ylabel("厚度平均压力 (MPa)")
axes[1, 0].set_ylabel("厚度平均孔隙率 (%)")
fig.suptitle("电极/隔膜局部压力—孔隙率径向演化", fontsize=14)
save(fig, "07_局部压力与孔隙率径向分布.png")


# 08 指标汇总图
fig, ax = plt.subplots(figsize=(11, 5.8))
ax.axis("off")
rows = [
    ("额定容量 / 0.5C电流", f"{metrics['rated_capacity_Ah']*1000:.3f} mAh / {metrics['applied_current_A']*1000:.3f} mA"),
    ("充电容量 / 放电容量", f"{metrics['charge_capacity_Ah']*1000:.3f} / {metrics['discharge_capacity_Ah']*1000:.3f} mAh"),
    ("库仑效率 / 能量效率", f"{metrics['coulombic_efficiency_pct']:.2f}% / {metrics['energy_efficiency_pct']:.2f}%"),
    ("充电截止 / 放电截止", f"{metrics['charge_end_time_s']/3600:.3f} h / {metrics['discharge_end_time_s']/3600:.3f} h"),
    ("最大 SOC / 最终 SOC", f"{metrics['maximum_actual_SOC_pct']:.2f}% / {metrics['final_actual_SOC_pct']:.2f}%"),
    ("60 s / 最终电压回弹", f"{metrics['rebound_60s_mV']:.1f} / {metrics['rebound_final_mV']:.1f} mV（相对最后放电输出点）"),
    ("石墨 / LFP 最大呼吸应变", f"{metrics['maximum_graphite_breathing_pct']:.2f}% / {metrics['maximum_lfp_breathing_abs_pct']:.2f}%"),
    ("平均压力峰值 / 最小孔隙率", f"{metrics['pressure_peak_MPa']:.3f} MPa / {metrics['minimum_porosity']:.4f}"),
    ("最大浓度跨度 / 最大接触压降", f"{metrics['maximum_electrolyte_span_mol_m3']:.1f} mol/m³ / {metrics['maximum_abs_contact_drop_mV']:.3f} mV"),
]
table = ax.table(cellText=rows, colLabels=["指标", "结果"], cellLoc="left", colLoc="left", loc="center", colWidths=[0.36, 0.58])
table.auto_set_font_size(False)
table.set_fontsize(11)
table.scale(1, 1.75)
for (row, col), cell in table.get_celld().items():
    if row == 0:
        cell.set_facecolor("#D9E8FB")
        cell.set_text_props(weight="bold")
    elif row % 2 == 0:
        cell.set_facecolor("#F5F7FA")
ax.set_title("扣电 V2 完整耦合循环关键指标", fontsize=16, pad=18)
save(fig, "08_关键指标汇总.png")


# Excel 汇总
book = Workbook()
ws = book.active
ws.title = "关键指标"
ws.append(["指标", "数值"])
for key, value in metrics.items():
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False)
    ws.append([key, value])
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5597")
    cell.alignment = Alignment(horizontal="center")
ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 28

ws2 = book.create_sheet("完整时序")
for row in [df.columns.tolist(), *df.itertuples(index=False, name=None)]:
    ws2.append(list(row))
for cell in ws2[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5597")
ws2.freeze_panes = "A2"

ws3 = book.create_sheet("SOC阶段")
headers = ["phase", "target_SOC", "time_s", "actual_SOC", "solnum"]
ws3.append(headers)
for stage in payload["stages"]:
    ws3.append([stage[h] for h in headers])
for cell in ws3[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5597")
for col in "ABCDE":
    ws3.column_dimensions[col].width = 18

book.save(OUT / "扣电V2_0p5C完整循环结果汇总.xlsx")
print(json.dumps({"ok": True, "figures": len(list(FIG.glob('*.png'))), "workbook": str(OUT / '扣电V2_0p5C完整循环结果汇总.xlsx')}, ensure_ascii=False))
