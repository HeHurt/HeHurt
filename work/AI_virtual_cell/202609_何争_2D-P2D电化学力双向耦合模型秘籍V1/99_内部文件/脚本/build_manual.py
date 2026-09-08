from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TEMPLATE = Path(r"C:\Users\hez\AppData\Local\Temp\codex-hithium\f413_2dp2d_manual\F413_template_plain.docx")
ASSETS = Path(r"C:\Users\hez\AppData\Local\Temp\codex-hithium\f413_2dp2d_manual\assets")
OUT = Path(r"C:\Users\hez\AppData\Local\Temp\codex-hithium\f413_2dp2d_manual\HC_TS_ISC_F413_2026_2D-P2D双向耦合模型.docx")

BLUE = "2F5597"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
RED = RGBColor(192, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF", size=9)
        set_cell_shading(table.rows[0].cells[i], BLUE)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=8.5)
            if ridx % 2:
                set_cell_shading(cells[i], LIGHT_GRAY)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="正文部分" if "正文部分" in [s.name for s in doc.styles] else "Normal")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="正文部分" if "正文部分" in [s.name for s in doc.styles] else "Normal")
    p.paragraph_format.left_indent = Inches(0.25 + 0.2 * level)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("■ " if level == 0 else "— ").bold = True
    p.add_run(text)
    return p


def add_equation(doc, equation, note=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(equation)
    r.font.name = "Cambria Math"
    r.font.size = Pt(11)
    if note:
        pn = doc.add_paragraph(note)
        pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for rr in pn.runs:
            rr.font.size = Pt(8.5)
            rr.font.color.rgb = RGBColor(89, 89, 89)


def add_figure(doc, filename, caption, width=6.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    c = doc.add_paragraph(style="图表" if "图表" in [s.name for s in doc.styles] else "Normal")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    r.font.size = Pt(9)
    return p


def add_note_box(doc, title, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_cell_shading(table.cell(0, 0), fill)
    p = table.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title + "：")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p.add_run(text)
    doc.add_paragraph()


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_page_break(doc):
    doc.add_page_break()


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "目录将在 Word 中更新"
    fld_sep.append(txt)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, fld_sep, fld_end])


doc = Document(str(TEMPLATE))
body = doc._element.body
sect_pr = body.sectPr
for child in list(body):
    if child is not sect_pr:
        body.remove(child)

# Preserve template page/header/footer system while making cover clean.
for section in doc.sections:
    section.different_first_page_header_footer = True
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

styles = doc.styles
styles["Normal"].font.name = "宋体"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Normal"].font.size = Pt(10.5)
for name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 11, "000000")]:
    st = styles[name]
    st.font.name = "黑体"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = True

doc.core_properties.title = "基于2D-P2D模型搭建的电化学-力双向耦合模型"
doc.core_properties.subject = "海辰仿真秘籍 HC_TS_ISC_F413_2026"
doc.core_properties.author = "何争"
doc.core_properties.keywords = "COMSOL, 2D-P2D, 电化学-力双向耦合, 压力, 孔隙率, 嵌锂呼吸"

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(75)
r = p.add_run("海辰仿真秘籍")
r.bold = True; r.font.name = "黑体"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体"); r.font.size = Pt(30); r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(25)
r = p.add_run("HC_TS_ISC_F413_2026")
r.bold = True; r.font.size = Pt(18); r.font.name = "Arial"
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(24)
r = p.add_run("《基于2D-P2D模型搭建的\n电化学-力双向耦合模型》")
r.bold = True; r.font.name = "黑体"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体"); r.font.size = Pt(24)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(95)
r = p.add_run("储能与管理研究院 · 电化学组")
r.font.size = Pt(14); r.bold = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("编制：何争        版本：V1.0        日期：2026-09").font.size = Pt(11)
add_page_break(doc)

# Compilation notes
add_heading(doc, "编制说明", 1)
add_body(doc, "本秘籍基于扣式电池与双侧出极耳单层软包电芯的 COMSOL 2D-P2D 建模、排错、力学耦合、参数扫描及后处理全过程编制。内容以已保存模型、求解数据、JSON 验收结果和正式图表为证据，覆盖从异常电压诊断到正确极性几何重建的完整演进。")
add_note_box(doc, "使用边界", "本文是一份可复现的内部建模方法，不是材料参数标准。弹性模量、泊松比、压力—孔隙率系数、接触电阻—压力关系及 SOC—膨胀曲线均需由对应材料/结构试验标定；未标定项在本文中按“经验值/待验证”管理。")
add_table(doc, ["版本", "日期", "编制/修订", "主要内容"], [
    ["V1.0", "2026-09", "何争", "汇总扣电与软包 2D-P2D 电化学—力双向耦合建模全流程、验证结果与问题闭环"],
])
add_heading(doc, "证据等级约定", 2)
add_bullet(doc, "A级—已在最终/当前模型中完成求解并由结果文件复核，可作为本轮仿真结论。")
add_bullet(doc, "B级—已在历史版本中完成单因素或机制验证，用于解释趋势，但不能替代最终几何复算。")
add_bullet(doc, "C级—经验参数、文献映射或待开展试验，只能作为建模输入建议。")
add_page_break(doc)
add_heading(doc, "目录", 1)
add_toc(doc)
add_page_break(doc)

# 1
add_heading(doc, "1 技术情况介绍", 1)
add_heading(doc, "1.1 建模目标", 2)
add_body(doc, "在传统 2D-P2D 电化学模型上引入力学场，使外部夹紧力和材料嵌锂呼吸共同改变局部压力、孔隙率、有效传输系数及接触电阻，并将这些变化反馈到电化学方程，从而得到电压、容量、极化、膨胀和应力的双向响应。")
add_bullet(doc, "对象一：CR2430 轴对称扣式电池，最终采用上部负极壳、下部正极壳的正确极性及参数化闭合堆叠。")
add_bullet(doc, "对象二：双侧出极耳单层软包 2D 截面，额定容量 0.13 Ah，0.5C 电流为 0.065 A；面外厚度必须与容量定义一致。")
add_bullet(doc, "统一协议：0.5C 充电—静置—0.5C 放电—静置，电压限值 3.65/2.50 V，总时长 18 000 s。")
add_heading(doc, "1.2 为什么需要双向耦合", 2)
add_body(doc, "只施加外力而不反馈电化学参数，电压几乎不会变化；只给定统一孔隙率，则无法反映局部压力不均。双向耦合的关键是让空间压力 p_local(r,z,t) 进入孔隙率 ε_l(r,z,t)、有效扩散/电导和接触电阻，同时让 SOC 控制的嵌锂应变返回 Solid Mechanics。")
add_figure(doc, "08_coupling_principle.png", "图1  嵌锂呼吸双向耦合的物理链路与 COMSOL 设置框架", 6.6)

# 2 theory
add_page_break(doc)
add_heading(doc, "2 2D-P2D 与力学耦合理论", 1)
add_heading(doc, "2.1 2D-P2D 电化学骨架", 2)
add_body(doc, "宏观二维几何求解集流体、电极和隔膜内的电势与电解液传输；每个多孔电极位置再附加一个颗粒半径维度求解固相锂扩散。因此模型既能得到平面内电流/压力不均，又保留 P2D 对颗粒动力学的描述。")
add_equation(doc, "∇·iₛ = −aₛj，   iₛ = −σₛ,eff∇φₛ")
add_equation(doc, "∂(εₗcₗ)/∂t = ∇·(Dₗ,eff∇cₗ) + (1−t₊⁰)aₛj/F")
add_equation(doc, "∂cₛ/∂t = (1/rₚ²)∂/∂rₚ(Dₛrₚ²∂cₛ/∂rₚ)")
add_equation(doc, "j = 2i₀sinh[Fη/(2RT)]，   η = φₛ−φₗ−U(θ)")
add_body(doc, "端电压取正极端子与负极接地端的电势差。额定容量决定 I₁C，几何面积决定边界平均电流密度，两者不可混用。")
add_heading(doc, "2.2 固体力学与嵌锂呼吸", 2)
add_equation(doc, "∇·σ = 0，   ε = εᵉ + εˢʷ，   σ = C:(ε−εˢʷ)")
add_equation(doc, "εˢʷ_k(θ_k) = f_k(θ_k)−f_k(θ_k,ref)，  k∈{Graphite, LFP}")
add_body(doc, "f_k 不应简单设为全 SOC 线性函数。石墨存在分级相变，膨胀曲线具有明显平台与斜率变化；LFP 变化幅度较小但同样可采用平滑非线性曲线。历史 20% 满量程线性石墨膨胀仅用于收敛性压力测试，不作为推荐材料参数。")
add_figure(doc, "05_nonlinear_breathing_curve.png", "图2  导入模型的典型非线性 SOC—嵌锂呼吸关系", 6.4)
add_heading(doc, "2.3 局部压力—孔隙率—传输反馈", 2)
add_equation(doc, "p_local(x,y,t) = max[0, −σ_yy(x,y,t)]")
add_body(doc, "对厚度方向受压的软包优先使用 −σyy；轴对称扣电按堆叠厚度方向选取相应法向应力。若材料试验对应三轴围压，可改用静水压力 p_h=max[0,−tr(σ)/3]。")
add_equation(doc, "εₗ,k = clip{εₗ,k,0[1−β_k p_local], εₗ,k,min, εₗ,k,max}")
add_equation(doc, "Dₗ,eff = Dₗ εₗᵇ，   κₗ,eff = κₗ εₗᵇ")
add_body(doc, "当前模型采用带上下限的经验压力反馈以保证数值稳定。更严格的有限应变表达可使用 εₗ=(εₗ,ref+ε_v)/(1+ε_v)，但需要材料体积分数、压缩曲线和变形框架一致性验证后再替换。")
add_heading(doc, "2.4 压力相关接触电阻", 2)
add_equation(doc, "R_c(p) = R_min + (R_0−R_min)exp(−p/p_ref)，   ΔV_contact = I·R_c")
add_body(doc, "接触电阻用端子/界面压降或等效 Rc 表征；浓差极化用电解液浓度极差 Δcₗ、空间梯度和电解液电势损失表征。二者必须分开诊断：前者在电流切换瞬间响应快，后者具有扩散时间尺度。")

# 3 diagnostics/version evolution
add_page_break(doc)
add_heading(doc, "3 扣电模型异常诊断与版本演进", 1)
add_heading(doc, "3.1 初始 2.95 V 平台并非单一电流问题", 2)
add_body(doc, "最初模型在约 2.95 V 附近变平并下跌。逐项检查表明，问题由多个设置叠加：1C 电流曾固定为 0.6 mA、活性面积/颗粒尺度不一致；多孔电极固相电导未正确生效，产生约 0.214 V 的异常集流体—电极压降；临时 V2 又使用最高约 3.55 V 的人工 LFP 曲线、反转正极嵌锂坐标并减去 20 mV，使 3.65 V 截止不可达。")
add_figure(doc, "01_ocv_curves.png", "图3  排错阶段正负极 OCV 曲线与全电池开路电压核查", 6.2)
add_figure(doc, "02_charge_curve_diagnosis.png", "图4  电导率、OCV 坐标和电流边界修正前后的充电曲线对比", 6.2)
add_note_box(doc, "结论", "原始 LFP.dat 与正极直接嵌锂坐标的思路是正确的；最初结果错误并不是因为“原始 OCV 一定错”，而是电流/面积、固相导电路径、临时人工 OCV 与坐标变换等多项设置共同造成。V4 恢复原始 LFP.dat、正极直接坐标及充电侧 +20 mV 后，OCV 主链恢复合理。")
add_heading(doc, "3.2 功能版本路线", 2)
add_table(doc, ["版本", "新增/修复", "验证状态", "主要用途"], [
    ["V2–V5", "容量面积、固相电导、原始 OCV、平均电流密度、电解液参数绑定", "历史验证(B)", "恢复合理电压主链"],
    ["V6", "Solid Mechanics 与外力扫描", "历史验证(B)", "单向压力敏感性"],
    ["V7", "局部压力→孔隙率→Deff/κeff 与 Rc(p)", "历史验证(B)", "空间非均匀反馈"],
    ["V8", "SOC 呼吸与压力反馈", "历史验证(B)", "线性呼吸/高膨胀收敛性"],
    ["V9", "文献型非线性呼吸与完整充放电事件", "已求解(B)", "双向耦合机制对照"],
    ["正确极性V2", "上负极壳/下正极壳、参数化闭合堆叠、探针重绑定", "最终求解(A)", "当前扣电基准"],
])
add_heading(doc, "3.3 探针与事件的隐藏风险", 2)
add_body(doc, "几何重建会改变边界编号并清除原边界探针。正确极性 V2 曾因电压探针仍指向失效边界而触发异常事件；将终端电压和加载端探针重新绑定到已验证边界后，完整循环恢复。任何参数化几何更新后都必须重新检查 named selections、端子、Ground、Fixed Constraint、Boundary Load 与事件表达式。")

# 4 modeling setup
add_page_break(doc)
add_heading(doc, "4 几何、参数与物理场设置", 1)
add_heading(doc, "4.1 扣电：正确极性与参数化闭合堆叠", 2)
add_bullet(doc, "二维轴对称：上部为负极壳/负端，下部为正极壳/正端；LFP、隔膜、石墨、铝箔、铜箔按真实堆叠顺序定义。")
add_bullet(doc, "几何采用厚度参数链。修改正/负极片厚度时，隔膜、集流体、壳体及密封结构的位置由累计厚度表达式平移，避免手工改坐标。")
add_bullet(doc, "壳体保持闭合，密封/绝缘域必须参与几何闭合检查；不能用开放缺口替代原结构。")
add_bullet(doc, "推荐 named selections：dom_lfp、dom_sep、dom_graphite、dom_al、dom_cu、bnd_neg_terminal、bnd_pos_terminal、bnd_load、bnd_fixed。")
add_heading(doc, "4.2 软包：2D 截面与面外厚度", 2)
add_bullet(doc, "电极区与集流体之间保留设计间距，极耳焊接在集流体顶端中心，而不是电极边缘。")
add_bullet(doc, "Y 轴放大 100 倍仅用于显示，不改变真实物理尺寸；力学厚度方向为 y，应使用 solid.sy 与位移 v。")
add_bullet(doc, "2D 模型必须设置正确的 out-of-plane thickness。若面外厚度与额定容量不一致，会导致 18 000 s 内无法完成一圈。")
add_heading(doc, "4.3 容量、电流和边界", 2)
add_equation(doc, "I₁C = Q_rated/1 h，   I₀.₅C = 0.5 I₁C，   j_app = I/A_terminal")
add_table(doc, ["模型", "额定容量", "0.5C 电流", "外力示例", "充放电窗口"], [
    ["扣电正确极性V2", "0.0163985 Ah", "8.1993 mA", "100 N", "3.65–2.50 V"],
    ["单层软包V8", "0.13 Ah", "65 mA", "0/100/250/500/750 N", "3.65–2.50 V"],
])
add_heading(doc, "4.4 网格简化原则", 2)
add_body(doc, "二维薄层模型不宜使用全局自动极细网格。对电极/隔膜厚度方向使用 mapped/swept 思路：隔膜至少 4–6 层单元，多孔电极各 6–10 层；径向/长度方向在极耳、端部、材料交界与载荷转折处局部加密，壳体远场可显著粗化。先做网格无关性：端电压、截止容量、最大 Δcₗ 和平均压力变化均小于 1% 后再固定。")
add_note_box(doc, "当前限制", "正确极性扣电 V2 的 1966 个单元可完成循环，但最小网格质量约 0.00369；因此局部峰值应力只作定性判断，平均压力、积分位移和电化学量更可靠。后续应优先改善尖角/薄缝网格质量。", "FFF2CC")

# 5 workflow
add_page_break(doc)
add_heading(doc, "5 COMSOL 实施流程", 1)
add_heading(doc, "5.1 建模顺序与验收闸门", 2)
add_table(doc, ["步骤", "实施内容", "必须通过的检查"], [
    ["1 几何", "参数化坐标、闭合堆叠、极性与极耳位置", "Build All 无错误；域数/相邻关系正确"],
    ["2 选区", "域、端子、加载面、固定面均用 named selections", "参数变更后选区仍非空"],
    ["3 电化学", "材料、OCV、粒径、容量、平均电流密度", "0 N 短时充电方向与起始 OCV 合理"],
    ["4 力学", "线弹性、固定约束、边界载荷", "Stationary/短时位移方向和量级合理"],
    ["5 单向反馈", "p_local→ε_l→Deff/κeff/Rc", "0/100 N 单点收敛且趋势可解释"],
    ["6 双向呼吸", "θ→εsw→σ→ε_l→电化学", "单工况完整循环先通过"],
    ["7 参数扫描", "外力/倍率/厚度", "各 case 事件完整、无探针丢失"],
    ["8 后处理", "阶段识别、匹配 SOC/截止状态", "指标定义一致，图与 JSON 可回读"],
])
add_heading(doc, "5.2 求解器与事件", 2)
add_bullet(doc, "推荐分步初始化：0 N 无呼吸短时 → 100 N 无呼吸 → 开启非线性呼吸 → 完整循环。")
add_bullet(doc, "Time Dependent 采用较小初始步，电流切换/截止附近限制最大步长；非线性强时使用 segregated 或适度阻尼，但不得用放宽容差掩盖物理错误。")
add_bullet(doc, "事件状态机至少包含充电、静置、放电、静置四阶段；电流必须由阶段变量唯一控制，截止电压必须引用真实端子电压。")
add_heading(doc, "5.3 后处理阶段抽取", 2)
add_body(doc, "优先按事件变量识别阶段，再在充电段按全局/电极平均 SOC 选择 10%、30%、50%、70%、85% 和截止时刻。对比呼吸前后或不同外力时，必须同时给出“同一时间”“同一 SOC”“各自截止”三种口径，避免把截止状态差异误判为材料机制。")
add_bullet(doc, "电压/性能：V(t)、V(SOC)、容量、库仑效率、能量效率、回弹 60 s/最终值。")
add_bullet(doc, "力学：平均/峰值厚度压力、von Mises、加载面位移、厚度变化。")
add_bullet(doc, "传输：平均/最小孔隙率、Δcₗ、κeff/Deff、接触压降。")
add_bullet(doc, "空间场：集流体 |J|、电极/隔膜界面 θ、正负极/集流体界面 θ、p_local 与 ε_l 径向/长度分布。")

# 6 results coin final
add_page_break(doc)
add_heading(doc, "6 扣电正确极性 V2：完整循环结果", 1)
add_heading(doc, "6.1 工况与循环完整性", 2)
add_figure(doc, "14_coin_v2_full_cycle.png", "图5  正确极性扣电 V2 的 0.5C 充电—静置—放电—静置完整循环", 6.4)
add_table(doc, ["指标", "结果", "证据等级"], [
    ["充电截止 / 放电开始 / 放电截止", "7115.72 / 7718.72 / 14706.99 s", "A"],
    ["充电 / 放电容量", "16.207 / 15.916 mAh", "A"],
    ["库仑效率 / 能量效率", "98.21% / 95.66%", "A"],
    ["最高平均 SOC / 末端 SOC", "99.829% / 2.77%", "A"],
    ["放电后回弹（60 s / 最终）", "344.4 / 422.7 mV（相对放电末点）", "A"],
])
add_figure(doc, "15_coin_v2_soc.png", "图6  全局 SOC 与正负极平均嵌锂状态演化", 6.4)
add_body(doc, "正负极嵌锂状态在充放电过程中呈相反方向是正常物理结果：充电时 Li 从 LFP 脱出并嵌入石墨，负极 θ 上升、正极 θ 下降；放电时方向相反。图中位移正负相反则来自坐标方向与约束定义，不代表一侧材料发生反常收缩。")
add_heading(doc, "6.2 力学与传输反馈", 2)
add_figure(doc, "16_coin_v2_mechanics.png", "图7  压力、孔隙率、非线性呼吸与加载面位移的时间响应", 6.4)
add_table(doc, ["量", "结果", "解释"], [
    ["石墨最大呼吸", "4.985%", "非线性 SOC 曲线在本工况实际达到值"],
    ["LFP 最大绝对呼吸", "0.598%", "量级明显低于石墨"],
    ["平均压力峰值", "0.313 MPa", "外力与嵌锂呼吸共同贡献"],
    ["最小孔隙率", "0.2916", "压力高区传输通道收缩"],
    ["加载边界位移范围", "259.77–297.76 μm", "相对变化约 38 μm"],
])
add_figure(doc, "17_coin_v2_polarization.png", "图8  电解液浓差极化与压力相关接触压降", 6.4)
add_body(doc, "本工况电解液浓度空间跨度最大约 909.7 mol·m⁻³，接触压降最大约 0.125 mV。由量级可见，当前电压差异主要不能归因于接触压降单项；需要结合浓差、反应不均和各自截止 SOC 综合判断。")

add_page_break(doc)
add_heading(doc, "6.3 空间分布：电流、嵌锂与局部压力", 2)
add_figure(doc, "18_coin_v2_current.png", "图9  不同 SOC 阶段集流构件电流密度分布", 6.55)
add_body(doc, "读图时首先使用统一色标并排除零电流静置段；重点比较壳体/集流体拐角、接触区域和有效电极边缘的热点位置。尖角极大值对网格敏感，建议同时报告 p95/p99 与面积平均值。")
add_figure(doc, "19_coin_v2_interface.png", "图10  正负极—隔膜及正负极—集流体界面的嵌锂状态分布", 6.55)
add_body(doc, "界面 θ 接近 0 或 1 表示局部活性材料先于平均 SOC 到达化学计量边界，可能触发提前截止或数值困难。应同时看平均值和径向分布，不能用单个极值代表整体利用率。")
add_figure(doc, "20_coin_v2_local_pressure.png", "图11  代表 SOC 下局部压力与孔隙率径向分布", 6.55)
add_body(doc, "压力高处孔隙率下降更多，进而降低 Deff 和 κeff。若整条孔隙率曲线随时间无波动，通常意味着使用了统一外压 p_ext、力学解未随时间更新，或 SOC 呼吸没有正确进入应变源。")

# 7 historic mechanism verification
add_page_break(doc)
add_heading(doc, "7 历史机制验证：从外力到非线性呼吸", 1)
add_heading(doc, "7.1 外力单向扫描", 2)
add_figure(doc, "03_force_mechanics.png", "图12  扣电 V6 外力扫描下的位移、压力与应力响应", 6.4)
add_table(doc, ["外力", "平均压力", "最大位移", "最大 von Mises", "SOC50 电压偏移"], [
    ["0 N", "0 MPa", "0 μm", "0 MPa", "0 mV"],
    ["100 N", "0.2137 MPa", "19.52 μm", "125.25 MPa", "+0.153 mV"],
    ["250 N", "0.5342 MPa", "48.81 μm", "313.13 MPa", "+0.565 mV"],
    ["500 N", "1.0684 MPa", "97.62 μm", "626.26 MPa", "+1.322 mV"],
])
add_body(doc, "该阶段用于验证力学载荷和接触电阻的单向趋势，不能直接等同于含呼吸的最终双向耦合结果。高应力峰值对尖角网格和经验材料参数敏感。")
add_heading(doc, "7.2 局部压力反馈", 2)
add_figure(doc, "04_local_pressure_porosity.png", "图13  扣电 V7 局部压力—孔隙率反馈及截止响应", 6.4)
add_body(doc, "100 N 时负极、隔膜、正极平均压力约为 0.206、0.210、0.215 MPa，局部最小孔隙率分别约 0.3026、0.3360、0.2760。引入空间孔隙率后，电解液浓度跨度从 V6 的约 544 mol·m⁻³ 增至 V7 的约 887 mol·m⁻³，说明局部反馈显著增强浓差非均匀性。")
add_heading(doc, "7.3 非线性呼吸替代 20% 线性试算", 2)
add_figure(doc, "06_breathing_voltage_cycle.png", "图14  关闭呼吸、20% 线性呼吸与文献型非线性呼吸的完整循环电压对比", 6.4)
add_figure(doc, "07_breathing_mechanics.png", "图15  不同呼吸机制下压力、孔隙率与位移多角度对比", 6.4)
add_body(doc, "历史 20% 线性石墨呼吸可收敛，但最大实际呼吸约 18.31%，远高于典型石墨电极的约 5% 量级。文献型非线性方案的最大石墨/LFP 呼吸约 4.985%/0.598%，压力峰值约 0.321 MPa，最小孔隙率约 0.2915，更适合作为当前工程基线。")

# 8 pouch and comparison
add_page_break(doc)
add_heading(doc, "8 单层软包迁移与扣电对比", 1)
add_heading(doc, "8.1 软包容量闭环", 2)
add_body(doc, "软包初次 18 000 s 未完成一圈的直接原因不是动力学参数，而是 2D out-of-plane thickness 与容量/电流的尺度不匹配。将额定容量明确为 0.13 Ah、I₁C=0.13 A、0.5C=0.065 A，并同步修正面外厚度后，可完成完整循环。")
add_heading(doc, "8.2 相同材料参数不等于相同电压", 2)
add_figure(doc, "09_coin_pouch_voltage.png", "图16  扣电与双侧出极耳软包的完整 0.5C 循环电压对比", 6.4)
add_body(doc, "两者复用了同一套 LFP/石墨 OCV、动力学和基础传输参数，但几何面积、面外厚度、集流路径、端子面积、外力作用面积和边界条件不同。因此总欧姆压降只是差异来源之一，反应电流密度分布、浓差极化、截止时局部 θ 和接触压力同样会改变曲线。")
add_figure(doc, "10_coin_pouch_metrics.png", "图17  扣电与软包容量、效率、回弹及非均匀性指标对比", 6.4)
add_figure(doc, "11_coin_pouch_mechanics.png", "图18  扣电与软包的压力、孔隙率、呼吸和厚度响应对比", 6.4)
add_heading(doc, "8.3 集流体热点", 2)
add_figure(doc, "12_pouch_current_maps.png", "图19  软包不同 SOC 阶段集流体/极耳电流密度分布", 6.55)
add_body(doc, "软包薄集流体和集中极耳使局部电流密度 p99 明显高于扣电分布式壳体路径。热点应优先在极耳焊接区、集流体过渡区和电极有效区边缘读取；若几何尖角未倒圆，峰值只作定位，不作绝对强度验收。")

add_page_break(doc)
add_heading(doc, "8.4 电压回弹的匹配诊断", 2)
add_figure(doc, "13_matched_rebound.png", "图20  软包 100 N 压力反馈基线与压力+呼吸的同协议匹配回弹", 6.4)
add_body(doc, "早期图上看似“扣电开启呼吸后回弹增大、软包反而减小”，实为基线、端点和回弹定义未完全匹配。按同一软包几何、100 N、相同 0.5C 协议和相同回弹定义重新提取后：压力-only 总回弹 386.586 mV，压力+呼吸为 391.604 mV，即增加 5.02 mV；瞬时回弹也由 62.984 增至 64.054 mV。")
add_note_box(doc, "回弹判读", "回弹 = 静置电压 − 放电末端电压。放电末端可能由电压截止、局部 θ 边界或事件触发决定。比较不同压力/呼吸时应同时报告截止 SOC、放电容量、末端 Δcₗ、Rc 与压力，不能仅比较静置平台高低。")
add_heading(doc, "8.5 外力扫描的非单调回弹", 2)
add_body(doc, "软包 0→250 N 回弹逐渐下降、250→750 N 又上升，可由两类竞争机制解释：中低压力下接触改善、Rc 降低使瞬时欧姆回弹减小；高压力下孔隙率下降、离子传输受限和截止状态改变，使浓差松弛及 OCV 恢复贡献增大。该解释属于模型机理假设，需用同截止 SOC/同放电容量重采样，并结合 EIS/DCR、厚度压力和浓度代理量试验验证。")

# 9 validation and FAQ
add_page_break(doc)
add_heading(doc, "9 结果验收、常见问题与试验标定", 1)
add_heading(doc, "9.1 最小验收清单", 2)
add_table(doc, ["类别", "验收项", "判定"], [
    ["几何", "极性、闭合、材料域、极耳/端子位置、参数化平移", "全部通过"],
    ["容量", "Q_rated、I1C、面外厚度、端子面积一致", "库仑量闭环"],
    ["OCV", "正负极坐标、单位、端点、温度和充放电支路", "无负电压/异常跳变"],
    ["事件", "充—静—放—静均出现，电流符号和截止正确", "时序完整"],
    ["力学", "载荷合力=设定外力，反力闭合，位移方向合理", "误差<1%"],
    ["耦合", "p_local 非负且时空变化；ε_l 不越界；呼吸随 SOC", "变量链不断裂"],
    ["网格", "V、容量、Δc_l、平均压力网格变化<1%", "独立性通过"],
    ["后处理", "同时间/同SOC/各自截止三口径不混用", "指标可追溯"],
])
add_heading(doc, "9.2 接触电阻与浓差极化如何试验表征", 2)
add_bullet(doc, "接触电阻：外力/夹紧压力扫描下的高频电阻、毫秒—1 s DCIR、四线法端子/界面电阻；结合卸载回线判断接触迟滞。")
add_bullet(doc, "浓差极化：不同倍率放电后的电压松弛、GITT/PITT、EIS 中低频扩散响应；模型侧读取 Δcₗ、∇cₗ 和电解液电势损失。")
add_bullet(doc, "孔隙率—压力：单层电极/隔膜厚度压缩、压汞/气体吸附或 XCT，至少覆盖装配压力范围并记录加载—卸载。")
add_bullet(doc, "SOC—呼吸：operando dilatometry、位移传感器或 3D 表面扫描；分别测试 LFP 与石墨电极，避免把整电芯热膨胀误当嵌锂呼吸。")
add_heading(doc, "9.3 常见问题定位表", 2)
add_table(doc, ["现象", "优先检查", "典型根因"], [
    ["起始电压过高", "OCV坐标、固相电导、端子压降", "正负极坐标反转或电导路径失效"],
    ["3.65 V不可达", "LFP OCV上限、事件、局部θ", "人工OCV上限过低或提前触边"],
    ["18000 s未完成循环", "容量、电流、面外厚度", "二维尺度不一致"],
    ["孔隙率无波动", "p_local时间依赖、呼吸应变源", "仍使用统一p_ext或力学未耦合"],
    ["压力分布像一根竖条", "坐标轴、显示缩放、选区", "软包厚度方向/绘图数据集选错"],
    ["结果突然掉压/报错", "探针、事件、局部θ和时间步", "几何更新后边界失效或化学计量越界"],
])

# 10 limitations/next
add_page_break(doc)
add_heading(doc, "10 适用边界与后续升级", 1)
add_heading(doc, "10.1 当前模型已包含", 2)
add_bullet(doc, "2D 宏观电流/浓度分布 + 颗粒径向扩散 P2D；")
add_bullet(doc, "Solid Mechanics、外部夹紧力、局部压力场；")
add_bullet(doc, "压力—孔隙率—有效传输和压力—接触电阻反馈；")
add_bullet(doc, "石墨/LFP 非线性 SOC—呼吸与双向反馈；")
add_bullet(doc, "完整充—静—放—静事件、外力扫描和多角度后处理。")
add_heading(doc, "10.2 尚未充分标定或未包含", 2)
add_bullet(doc, "电极、隔膜、密封件和壳体的真实非线性/粘弹塑性、加载卸载迟滞与长期蠕变；")
add_bullet(doc, "接触开闭、摩擦、脱层和焊点非线性；当前 Rc(p) 为等效经验关系；")
add_bullet(doc, "温度场、产热和热膨胀；当前回弹分析未拆分热松弛；")
add_bullet(doc, "SEI、锂析出、颗粒开裂等老化机制及循环累积残余应变；")
add_bullet(doc, "压力—孔隙率和 SOC—膨胀曲线的本项目材料实测标定。")
add_page_break(doc)
add_heading(doc, "10.3 推荐 DOE", 2)
add_table(doc, ["因子", "水平", "输出", "目的"], [
    ["外力", "0/100/250/500/750 N", "V、回弹、Rc、Δc_l、ε_l、压力", "识别接触改善与传输受限拐点"],
    ["倍率", "0.2C/0.5C/1C", "截止SOC、极化分解、热点", "区分欧姆与扩散贡献"],
    ["呼吸", "关闭/文献曲线/实测曲线", "位移、压力、回弹、效率", "校准双向耦合增量"],
    ["边界", "力控/位移控", "反力、厚度、孔隙率", "贴合实际夹具"],
    ["网格", "粗/中/细", "V、容量、平均压力、p99电流", "确认网格独立性"],
])

# 11 conclusions
add_page_break(doc)
add_heading(doc, "11 仿真结论", 1)
add_body(doc, "① 针对宇啸提供的扣电模型完成了从电压异常到正确极性重建的闭环。先前约 2.95 V 平台及提前掉压不是单一电流大小造成，而是容量/面积、固相导电路径、临时人工 LFP OCV 与嵌锂坐标、事件探针等设置叠加。通过恢复原始 LFP.dat、修正 OCV 坐标与固相电导、采用平均电流密度、重绑定事件探针，并最终按上负极壳/下正极壳重建参数化闭合堆叠，得到可完成 0.5C 完整循环的当前模型。")
add_body(doc, "② 嵌锂状态分析表明，充电时石墨嵌锂比例上升、LFP 嵌锂比例下降，放电反向；界面分布可识别局部先触及 θ=0/1 的提前截止风险。正确极性 V2 的最高平均 SOC 达 99.829%，充/放电容量为 16.207/15.916 mAh，库仑效率 98.21%。")
add_body(doc, "③ 过流部分电流密度分析表明，热点主要位于集流路径收缩、极耳/壳体转折及有效电极边缘。软包集中极耳路径使局部电流密度显著高于扣电分布式壳体，说明两种电芯即使材料动力学相同，也会因几何和边界产生不同电压与极化。")
add_body(doc, "④ 双向耦合结果显示，SOC 呼吸通过压力改变局部孔隙率和传输，再反馈端电压。正确极性扣电 V2 的石墨/LFP 最大呼吸约 4.985%/0.598%，平均压力峰值约 0.313 MPa，最小孔隙率约 0.2916。软包同工况匹配后，开启呼吸使总回弹增加约 5.02 mV，纠正了早期因对比口径不一致造成的相反判断。")
add_note_box(doc, "最终判断", "当前模型已具备工程机理研究所需的电化学—力双向耦合主链，但仍属于“结构与机制可用、材料参数待标定”的阶段。用于绝对寿命/安全判定前，必须完成压力—孔隙率、Rc(p)、非线性力学和 SOC—膨胀曲线的实测标定及网格独立性。")
add_figure(doc, "21_coin_v2_summary.png", "图21  正确极性扣电 V2 关键指标汇总", 6.4)

# References
add_page_break(doc)
add_heading(doc, "参考资料", 1)
refs = [
    "[1] COMSOL Multiphysics 6.4, Battery Design Module User’s Guide: Lithium-Ion Battery Interface and Heterogeneous Lithium-Ion Battery.",
    "[2] Rieger B., Schlueter S., Erhard S. V., et al. Multi-scale investigation of thickness changes in a commercial pouch type lithium-ion battery. Journal of Energy Storage, 2016, 6: 213–221. DOI: 10.1016/j.est.2016.01.006.",
    "[3] Escher I., et al. A Practical Guide for Using Electrochemical Dilatometry as Operando Tool in Battery and Supercapacitor Research. Energy Technology, 2022, 10: 2101120. DOI: 10.1002/ente.202101120.",
    "[4] 本项目扣电 V2–V9 与正确极性 V2 COMSOL 模型、求解结果、后处理 JSON/PNG（2026-08—2026-09）。",
    "[5] 本项目双侧出极耳单层软包 V4–V8 COMSOL 模型、0.13 Ah/0.5C 完整循环及外力扫描结果（2026-08—2026-09）。",
]
for ref in refs:
    add_body(doc, ref)

# Final method summary in the style of the source template.
add_page_break(doc)
add_heading(doc, "方法说明", 1)
add_heading(doc, "编制原因", 2)
add_body(doc, "将扣电与软包电芯从电压异常排查、2D-P2D 几何迁移、外力加载、局部压力反馈到非线性嵌锂呼吸的多轮工作沉淀为统一方法，避免后续模型重复踩容量尺度、极性、OCV 坐标、探针失效和对比口径不一致等问题。")
add_heading(doc, "主要内容", 2)
add_body(doc, "本文给出模型方程、COMSOL 设置顺序、参数化几何原则、网格与求解策略、完整循环事件、后处理指标、已验证结果、误判纠正、试验标定清单和后续升级边界。执行时以“先 0 N 基线、再单向压力、最后双向呼吸；先单工况、后参数扫描”为总原则。")

doc.settings.update_fields_on_open = True
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(OUT)
