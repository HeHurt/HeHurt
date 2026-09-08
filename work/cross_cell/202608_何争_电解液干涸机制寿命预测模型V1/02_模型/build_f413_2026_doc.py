"""Build HC_TS_ISC_F413_2026 from the retained department-manual reference."""

from pathlib import Path
import copy
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[4]
TASK = ROOT / "work" / "cross_cell" / "202608_何争_电解液干涸机制寿命预测模型V1"
REFERENCE = Path(r"E:\Downloads\HC_TS_ISC_F413_2025《基于PyBaMM计算电芯不同SOH下能效、功率性能、DCR的方法制定》.docx")
FINAL = TASK / "01_仿真报告" / "HC_TS_ISC_F413_2026《基于PyBaMM添加电解液干涸机制的寿命预测模型》.docx"
FIG = TASK / "04_输出结果" / "figures"


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=12, bold=None, color=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_margins(cell, top=60, start=100, bottom=60, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total_dxa = int(sum(widths_cm) / 2.54 * 1440)
    tblW.set(qn("w:w"), str(total_dxa))
    tblW.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width / 2.54 * 1440)))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = Cm(widths_cm[idx])
            cell.width = width
            tcW = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tcW.set(qn("w:w"), str(int(widths_cm[idx] / 2.54 * 1440)))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def shade_cell(cell, fill):
    shd = cell._tc.get_or_add_tcPr().first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        cell._tc.get_or_add_tcPr().append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_bottom_border(paragraph, color="707070", size="6"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "打开文档后更新目录"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_separate, fallback, fld_char_end])


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for name, size, before, after in (
        ("Heading 1", 18, 16, 10),
        ("Heading 2", 15, 13, 8),
        ("Heading 3", 13, 10, 6),
        ("Heading 4", 12, 8, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Pt(0)

    for style_name in ("List Number", "List Bullet"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            style.font.size = Pt(12)
            style.paragraph_format.space_after = Pt(4)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=12, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=12)
    else:
        r = p.add_run(text)
        set_run_font(r, size=12)
    return p


def get_numbering_id(doc, numbered):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if numbered else "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if numbered else "•")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list(doc, items, numbered=True):
    num_id = get_numbering_id(doc, numbered)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num])
        p_pr.append(num_pr)
        r = p.add_run(item)
        set_run_font(r, size=12)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run(text)
    set_run_font(r, east_asia="黑体", latin="Arial", size={1: 18, 2: 15, 3: 13, 4: 12}[level], bold=True)
    return p


def add_code(doc, lines):
    for line in lines.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.right_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F2F2F2")
        p._p.get_or_add_pPr().append(shd)
        r = p.add_run(line or " ")
        set_run_font(r, east_asia="等线", latin="Consolas", size=9)


def add_figure(doc, path, caption, width=Inches(5.75)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    cap.paragraph_format.keep_together = True
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, size=10.5)


def add_table(doc, headers, rows, widths_cm, caption=None, header_fill="D9E2F3"):
    if caption:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(caption)
        set_run_font(r, size=10.5)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths_cm)
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(text)
        set_run_font(r, east_asia="黑体", latin="Arial", size=10.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(text) > 16 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, size=9.0 if len(text) > 18 else 9.5)
    set_table_widths(table, widths_cm)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def build():
    doc = Document(str(REFERENCE))
    clear_body(doc)
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.first_line_indent = Pt(0)
    hp.text = ""
    hr = hp.add_run("基于PyBaMM添加电解液干涸机制的寿命预测模型")
    set_run_font(hr, size=8, color=(90, 90, 90))
    set_paragraph_bottom_border(hp, color="808080", size="5")
    section.first_page_header.paragraphs[0].text = ""

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Pt(0)
    add_field(fp, " PAGE ")
    section.first_page_footer.paragraphs[0].text = ""

    # Cover
    rule = doc.add_paragraph()
    rule.paragraph_format.first_line_indent = Pt(0)
    set_paragraph_bottom_border(rule, color="606060", size="8")
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    set_run_font(p.add_run("编号：HC_TS_ISC_F413_2026"), east_asia="黑体", latin="Arial", size=12, bold=True)
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    set_run_font(p.add_run("《基于PyBaMM添加电解液干涸机制的\n寿命预测模型》"), east_asia="黑体", latin="Arial", size=22, bold=True)
    for _ in range(3):
        doc.add_paragraph()
    for text in (
        "编  制：  何 争     日期： 2026-08-03",
        "校  对：  待确认     日期： __________",
        "审  核：  待确认     日期： __________",
        "批  准：  待确认     日期： __________",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(16)
        set_run_font(p.add_run(text), size=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    set_run_font(p.add_run("厦门海辰储能科技股份有限公司"), east_asia="宋体", size=12)
    page_break(doc)

    # Compilation note
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    set_run_font(p.add_run("编制说明"), east_asia="黑体", latin="Arial", size=18, bold=True)
    add_body(doc, "为使本公司基于PyBaMM开展电解液干涸寿命预测的工作规范化，结合BatteryProject现有代码实现、测试与项目应用经验，编制本仿真分析规范。")
    add_body(doc, "本规范用于说明电解液干涸机制的物理假设、参数入口、分块耦合流程、开关方式、后处理指标及验证要求，供仿真计算人员建立和审查寿命模型时使用。")
    add_body(doc, "本规范描述的是当前BatteryProject中的准静态体积守恒实现。凡涉及注液量、干涸速率、裂纹SEI或寿命精度的定量结论，均须结合目标电芯实验重新标定，不得将本文中的降阶核算示例直接作为产品预测结果。")
    add_body(doc, "本规范主要起草人：何争。")
    add_body(doc, "本规范的版本记录和版本号变动与修订记录如下。")
    add_table(
        doc,
        ["版本号", "制定/修订者", "制定/修订日期", "批准", "日期"],
        [
            ["HC_TS_ISC_F413_2026", "何争", "2026.08.03", "待确认", "待确认"],
            ["", "", "", "", ""],
            ["", "", "", "", ""],
            ["", "", "", "", ""],
        ],
        [4.7, 2.3, 2.9, 1.9, 1.9],
    )
    page_break(doc)

    # TOC
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    set_run_font(p.add_run("目录"), east_asia="黑体", latin="Arial", size=18, bold=True)
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Pt(0)
    add_field(toc, ' TOC \\o "1-3" \\h \\z \\u ')
    page_break(doc)

    # 1
    add_heading(doc, "1 技术情况介绍", 1)
    add_body(doc, "本规范规定了基于PyBaMM为电芯寿命预测模型增加电解液干涸机制的方法、输入、计算流程、后处理及验收要求。")
    add_body(doc, "本规范适用于以DFN/P2D为基础、且已建立SEI等副反应模型的锂离子电芯。对钠离子体系、无SEI溶剂消耗路径的模型或需要三维局部润湿分布的任务，须另行评估。")
    add_heading(doc, "1.1 仿真工具介绍", 2)
    add_body(doc, "PyBaMM（Python Battery Mathematical Modelling）是开源电池建模框架。BatteryProject在其基础上统一电芯参数、寿命工况、结果提取和项目交付。本规范使用的核心模块为BatteryProject/src/electrolyte_dryout.py。")
    add_list(doc, [
        "基础模型：DFN/P2D为主，SPM可用于低成本诊断，但应重新确认状态变量是否齐全。",
        "老化来源：负极SEI与裂纹SEI的锂损失量、三域平均孔隙率和电解液锂量。",
        "耦合方式：寿命按block分段求解，block之间更新电解液体积、浓度和等效有效宽度。",
        "求解器：优先IDAKLUSolver；寿命末期发生非物理状态或NaN时保留已完成block并停止。",
    ], numbered=False)
    add_heading(doc, "1.2 适用范围与模型定位", 2)
    add_body(doc, "当前实现不是在单个时间步内连续求解液体迁移的PDE子模型，而是一个块间准静态coupler。其价值是把SEI耗液、孔隙变化和储液余量转换为下一段PyBaMM仿真的有效初始条件，计算成本低、便于与现有寿命模型组合。")
    add_body(doc, "因此模型可回答“电解液余量何时不足、欠液后有效反应面积如何变化、干涸对容量/DCR/极化趋势的增量影响”等工程问题，但不能直接给出电芯内部空间干区位置、局部润湿前沿或气液两相流分布。")
    page_break(doc)

    # 2
    add_heading(doc, "2 仿真模型介绍", 1)
    add_heading(doc, "2.1 基础电化学与老化模型", 2)
    add_body(doc, "基础模型建议采用PyBaMM DFN，电化学参数先通过倍率、高低温和OCV等数据完成BOL校准；老化模型至少包含SEI，若启用裂纹SEI则应同时核对裂纹相关状态量。干涸参数不应承担补偿未校准动力学参数的作用。")
    add_list(doc, [
        "仅考虑极片厚度方向的主传输过程，颗粒采用球形径向扩散。",
        "温度可作为工况输入；未建立热耦合时，默认每个case温度恒定。",
        "干涸通过卷芯整体平均量耦合，不区分沿高度、宽度或厚度方向的局部差异。",
        "当前代码把欠液后的有效面积损失等效为Electrode width缩减。",
    ])
    add_heading(doc, "2.2 电解液干涸机理", 2)
    add_body(doc, "核心因果链为：SEI生成消耗EC溶剂，使全电芯电解液总体积下降；副反应产物或力学压缩改变电极/隔膜孔隙体积；储液区在孔隙与总液量之间承担补液或接收挤出液体；当储液区不足以填满卷芯孔隙时，Ratio_Dryout低于1，并通过等效宽度和浓度修正影响下一block。")
    add_figure(doc, FIG / "fig_2_1_dryout_mechanism.png", "图2-1 BatteryProject电解液干涸准静态耦合链")
    add_heading(doc, "2.3 关键状态量与可观测输出", 2)
    add_table(
        doc,
        ["类别", "代码状态量", "物理含义", "建议观测/校验"],
        [
            ["库存", "Vol_Elely_Tot", "全电芯剩余电解液体积", "注液量、失液量、拆解称重"],
            ["润湿", "Vol_Elely_JR / Vol_Pore_tot", "卷芯内液量与可用孔隙体积", "EOL浸润、局部干区证据"],
            ["干涸", "Ratio_Dryout", "卷芯欠液比例；小于1表示欠液", "容量拐点、DCR/极化突增"],
            ["浓度", "Ratio_CeEC / Ratio_CeLi", "EC和Li+浓度初值修正", "电解液分析、浓差极化"],
            ["等效面积", "Width / Width_0", "欠液导致的等效有效宽度", "容量、功率、DCR"],
        ],
        [2.0, 3.5, 4.5, 4.0],
        caption="表2-1 干涸追踪状态量及工程解释",
    )
    page_break(doc)

    # 3
    add_heading(doc, "3 仿真工况及假设条件", 1)
    add_body(doc, "干涸寿命预测必须先明确目标工况和是否具备可辨识的数据。相同容量保持率可由SEI、LAM、析锂、孔隙堵塞与干涸的不同组合产生，仅拟合容量曲线无法证明干涸机理正确。")
    add_list(doc, [
        "循环工况：按项目实际恒流/恒功率工况配置，明确SOC窗口、温度、静置时间与循环间RPT。",
        "分块策略：n_blocks × cycles_per_block覆盖目标寿命；block越小，干涸更新越细，阶梯现象越弱，但计算成本越高。",
        "加速因子：t_factor仅用于老化参数加速，不等于物理时间缩放；高精度或强非线性阶段应减小或关闭。",
        "初始电解液：excess_ratio应由注液量、卷芯孔隙体积和化成后有效余量换算，不得直接照搬其他电芯。",
        "温度切换：每个温度重新创建pybamm.ParameterValues(\"OKane2022\")并更新对应电芯参数，禁止复用被修改的ParameterValues。",
    ])
    add_body(doc, "建议至少设置“干涸关闭”和“干涸开启”两个同工况case。两者除tracker外保持模型、参数、求解器、网格和实验完全一致，用于识别干涸的增量贡献。")
    page_break(doc)

    # 4
    add_heading(doc, "4 设计/校对输入", 1)
    add_body(doc, "开展干涸寿命预测前，应由项目组和仿真人员共同确认表4-1中的输入。未提供的数据必须标记为待确认，并在结果中给出敏感性范围。")
    add_table(
        doc,
        ["序号", "输入项目", "参数/数据", "归属", "备注"],
        [
            ["1", "电芯设计与极片几何", "参数", "电化学模型", "正负极/隔膜厚度、宽度、高度、初始孔隙率"],
            ["2", "电解液注液与化成后余量", "参数/实测", "干涸模型", "用于换算excess_ratio；优先采用有效液量而非标称注液量"],
            ["3", "EC初始浓度与偏摩尔体积", "参数", "干涸模型", "Bulk solvent concentration、EC partial molar volume"],
            ["4", "短期寿命与容量保持率", "实测", "老化模型", "用于校准SEI/LAM等基线，不足以单独辨识干涸"],
            ["5", "DCR/功率/循环电压", "实测", "联合校验", "检查欠液后极化和电阻变化是否合理"],
            ["6", "EOL拆解/浸润/残液证据", "实测", "机理校验", "用于判断干涸是否真实发生及发生阶段"],
            ["7", "裂纹与析锂相关证据", "实测/诊断", "竞争机理", "避免把裂纹SEI或析锂导致的拐点误归因于干涸"],
        ],
        [1.2, 3.3, 2.3, 2.6, 4.7],
        caption="表4-1 电解液干涸寿命模型输入清单",
    )
    add_body(doc, "当前实现中，初始卷芯液量由几何孔隙体积计算，总液量为孔隙体积乘以excess_ratio。参数文件中历史保留的Initial total electrolyte volume in whole cell、Initial total electrolyte volume in jelly roll和Electrolyte dry out rate并非DryoutTracker主链当前直接使用的控制量，调参时应以代码实际读取项为准。")

    # 5
    add_heading(doc, "5 流程及方法", 1)
    add_body(doc, "标准流程为：基线电化学校准 → 基线老化校准 → 注液/孔隙参数核对 → 干涸tracker初始化 → 分块寿命求解 → 干涸更新 → 末态重启 → 干涸开关对比 → 多观测量验收。")
    add_heading(doc, "5.1 软硬件要求", 2)
    add_heading(doc, "5.1.1 软件要求", 3)
    add_list(doc, [
        "Python 3.10+，BatteryProject当前依赖环境。",
        "PyBaMM及IDAKLUSolver；CasadiSolver仅作为兼容性fallback。",
        "Visual Studio Code或Codex用于代码与Notebook操作。",
        "matplotlib + scienceplots用于后处理，字体采用Calibri + Microsoft YaHei。",
    ])
    add_heading(doc, "5.1.2 硬件要求", 3)
    add_body(doc, "最小case可在本地工作站运行；多温度、多设计、多敏感性矩阵建议使用多进程或服务器。首次验证必须从单温度、单工况、少量block开始，不得直接提交全矩阵。")
    add_heading(doc, "5.2 干涸参数设置流程", 2)
    add_heading(doc, "5.2.1 建立基线模型", 3)
    add_body(doc, "先完成无干涸DFN模型校准，并确认SEI、LAM、析锂、孔隙率等内部状态没有明显非物理趋势。干涸模型只应解释基线模型无法解释且有实验依据的欠液效应。")
    add_code(doc, '''model = pybamm.lithium_ion.DFN({
    "SEI": "reaction limited",
    "SEI porosity change": "true",
    "SEI on cracks": "true",          # 有裂纹证据时启用
    "loss of active material": "stress-driven",
    "lithium plating": "partially reversible",
})
solver = pybamm.IDAKLUSolver()
var_pts = {"x_n": 5, "x_s": 5, "x_p": 5, "r_n": 20, "r_p": 20}''')
    add_heading(doc, "5.2.2 初始化电解液库存", 3)
    add_body(doc, "DryoutTracker首先根据正极、负极、隔膜厚度与孔隙率，以及电极宽度、高度计算卷芯孔隙体积：")
    add_body(doc, "V_pore,0 = (L_n ε_n + L_p ε_p + L_s ε_s) × L_y × L_z")
    add_body(doc, "初始卷芯液量取V_pore,0，初始全电芯总液量取V_pore,0 × excess_ratio，二者差值作为等效储液区余量。excess_ratio≥1时tracker启用；工程使用建议从1.0及以上的实测换算值开始。")
    add_figure(doc, FIG / "fig_5_1_initial_inventory.png", "图5-1 excess_ratio对初始电解液库存的控制（由当前参数文件计算）")
    add_heading(doc, "5.2.3 核心体积守恒方程", 3)
    add_body(doc, "每个block结束后，代码读取负极SEI和裂纹SEI在该block内的锂损失增量，并用EC偏摩尔体积换算溶剂消耗：")
    add_body(doc, "V_EC,consumed = (Δn_LLI,SEI + Δn_LLI,SEI-crack) × V̄_EC")
    add_body(doc, "三域末态平均孔隙率用于重算孔隙体积。需要从储液区补入卷芯的体积定义为：")
    add_body(doc, "V_need = V_EC,consumed − (V_JR,old − V_pore,new)")
    add_body(doc, "当孔隙体积下降量大于溶剂消耗时，V_need<0，液体被挤入储液区；当V_need≥0时，优先从储液区补液；储备不足时形成Ratio_Dryout<1。")
    add_figure(doc, FIG / "fig_5_2_branch_logic.png", "图5-2 干涸体积平衡三类分支")
    add_heading(doc, "5.2.4 浓度与等效面积更新", 3)
    add_body(doc, "储液区参与补液时，代码更新卷芯内Li+总量与EC总量，并得到Ratio_CeLi_JR和Ratio_CeEC_JR。下一block建立初始条件时，对负极、隔膜、正极的“porosity times concentration”状态乘以Li+浓度修正比例。")
    add_body(doc, "储备不足时，Electrode width更新为Ratio_Dryout × L_y，用于等效表达润湿反应面积下降。由此定义干涸等效LAM：LAM_dryout = 100% − Width/Width_0 × 100%。该量是工程等效指标，不等同于材料真实失活。")
    add_body(doc, "若同时启用swelling_coupler，孔隙率压缩因子还会同步乘到ε·c_e状态，以保持浓度一致；被挤出的电解液在下一次DryoutTracker.update中进入储液区核算。")
    add_heading(doc, "5.2.5 分块寿命运行与开关", 3)
    add_code(doc, '''ENABLE_DRYOUT = True
tracker = DryoutTracker(params, excess_ratio=EXCESS_RATIO) if ENABLE_DRYOUT else None

sol_list = run_aging_with_dryout(
    model=model,
    params=params,
    experiment=build_cycle_experiment,
    solver=solver,
    var_pts=var_pts,
    tracker=tracker,
    n_blocks=N_BLOCKS,
    cycles_per_block=CYCLES_PER_BLOCK,
    t_factor=T_FACTOR,
    temperature=TEMPERATURE,
    get_hithium_params=get_hithium_params,
)''')
    add_body(doc, "关闭干涸的标准方式是tracker=None；开启方式是传入DryoutTracker实例。启用tracker或swelling_coupler后，get_hithium_params仅在首个block刷新，以避免覆盖累计变化的Electrode width等状态。")
    add_body(doc, "分块末态由apply_dryout_to_initial_conditions写入下一block。若寿命末期IDAKLU返回NaN/Inf或一致初值失败，run_aging_with_dryout捕获pybamm.SolverError，停止后续block并返回已完成结果。该行为用于保留证据，不代表失效点已通过物理验证。")
    add_heading(doc, "5.2.6 阶梯现象与数值收敛", 3)
    add_body(doc, "容量保持率、Ratio_Dryout或等效宽度出现阶梯，通常来自block之间离散更新，而不是单独的代码错误。减小cycles_per_block或增加n_blocks可使更新更细，但应检查总循环数不变，并进行block-size收敛性比较。")
    add_list(doc, [
        "先跑tracker=None基线，确认模型在目标循环范围内可解。",
        "再用excess_ratio较高的弱干涸case，确认开启耦合不会立即产生非物理状态。",
        "逐步降低excess_ratio或增强SEI耗液，不得一次跨越到强干涸。",
        "比较不同block大小下的干涸起点、EOL和关键指标；差异过大时不得发布。",
    ])
    add_heading(doc, "5.3 参数校准与可辨识性", 2)
    add_body(doc, "推荐按“先基线、后干涸”的顺序校准。首先用容量、电压和DCR约束SEI/LAM/析锂；随后用有效注液量、残液或拆解润湿证据约束excess_ratio；最后用干涸开启/关闭的差值解释容量、负极电位、平台位移、DCR和热量等可观测变化。")
    page_break(doc)
    add_table(
        doc,
        ["参数/假设", "主要影响", "不可单独依赖的观测", "优先约束证据"],
        [
            ["excess_ratio", "干涸起始block", "容量保持率", "有效注液量、残液、拆解润湿"],
            ["EC偏摩尔体积/消耗计量", "耗液速率", "单点EOL", "电解液组成、SEI副产物假设"],
            ["SEI/裂纹SEI动力学", "EC消耗来源", "容量衰减", "LLI分解、裂纹/阻抗诊断"],
            ["block大小", "阶梯粒度、数值误差", "单一离散曲线", "block-size收敛"],
            ["等效宽度假设", "容量/功率/DCR增量", "仅容量拟合", "多指标联合验证"],
        ],
        [3.0, 3.3, 3.5, 4.3],
        caption="表5-1 干涸参数辨识与证据要求",
    )
    # 6
    add_heading(doc, "6 后处理及功能展示", 1)
    add_body(doc, "DryoutTracker.history记录总液量、卷芯液量、孔隙体积、干涸比例、浓度比例、有效宽度、EC消耗、补液量、孔隙减少量和储液区浓度。优先调用tracker.plot()生成六联图，再按项目需要导出单项趋势。")
    add_heading(doc, "6.1 体积库存与干涸比例", 2)
    add_body(doc, "首先检查Vol_Elely_Tot是否随EC消耗单调下降；再检查Vol_Elely_JR与Vol_Pore_tot的关系。Ratio_Dryout=1表示储液区仍能填满卷芯孔隙，Ratio_Dryout<1表示欠液。V_need<0时模型允许液体被挤入储液区，此时Ratio_Dryout可高于1，应结合分支解释而非直接按“湿润率”裁剪。")
    add_heading(doc, "6.2 等效宽度与容量保持率", 2)
    add_body(doc, "欠液后Width/Width_0下降，等效反应面积和名义容量随之下降。建议将容量保持率、干涸等效LAM、总LLI、LAM、DCR和负极电位放在同一cycle/block坐标下，判断容量拐点是否与干涸起点同步。")
    add_figure(doc, FIG / "fig_6_1_reduced_order_demo.png", "图6-1 当前体积守恒代码的降阶分块核算示例（非寿命标定结果）")
    add_body(doc, "图6-1使用MIC参数文件和受控block末态调用当前DryoutTracker.update，目的是验证库存阈值和分支行为。结果表明，提高excess_ratio会推迟Ratio_Dryout低于1及等效宽度下降的开始位置。该图未运行完整DFN寿命仿真，不用于宣称真实循环寿命。")
    add_heading(doc, "6.3 干涸开启/关闭对比", 2)
    add_body(doc, "正式项目应至少输出同工况的Dryout OFF与Dryout ON曲线，并在独立表格中给出容量RMSE/RRMSE、干涸起点、EOL、DCR增量和求解终止block。测量曲线使用虚线，模型曲线使用实线；指标放入表格，不在图中堆叠。")
    page_break(doc)
    add_table(
        doc,
        ["输出项", "Dryout OFF", "Dryout ON", "实验/验收"],
        [
            ["容量保持率曲线", "待运行", "待运行", "待项目数据确认"],
            ["Ratio_Dryout起点", "不适用", "待运行", "EOL拆解/残液"],
            ["50%SOC DCR", "待运行", "待运行", "循环RPT"],
            ["负极电位/析锂风险", "待运行", "待运行", "三电极或机理判断"],
            ["求解终止block", "待运行", "待运行", "数值稳定性记录"],
        ],
        [3.5, 3.0, 3.0, 4.0],
        caption="表6-1 正式项目干涸开关对比结果表（待补充实算数据）",
        header_fill="E7E6E6",
    )
    add_body(doc, "表6-1保留为实算结果入口。当前任务没有指定目标电芯、寿命工况和实验对标数据，因此不填入虚构的容量/DCR数值；由后续项目case运行后替换“待运行”。")
    add_heading(doc, "6.4 代码级审查注意事项", 2)
    add_list(doc, [
        "体积消耗同时计入SEI与裂纹SEI的LLI；当前EC摩尔数更新式只显式扣减负极SEI项，强裂纹SEI场景需专项核对摩尔守恒。",
        "Ratio_Dryout和Electrode width未设置工程下限；出现接近零或负值前应提前终止，并记录最后一个有效block。",
        "现有单元测试覆盖初始化、历史字段和孔隙率因子状态传递；真实update与完整寿命case仍需要集成测试。",
        "参数文件中的历史电解液体积/干涸速率字段与当前tracker初始化逻辑并存，发布前应防止调用者误用。",
    ])
    page_break(doc)

    # 7
    add_heading(doc, "7 项目交付与反馈答疑", 1)
    add_body(doc, "完成计算后，应将任务说明、模型/脚本、输入数据、原始run、发布结果、图表和复现说明一并归档。正式返修创建V2/V3，不覆盖已上传版本。")
    add_heading(doc, "7.1 最小验收清单", 2)
    add_list(doc, [
        "基线无干涸case已通过BOL与短期寿命校准；干涸不是用于掩盖基线误差。",
        "excess_ratio有目标电芯依据或明确敏感性区间，不照搬其他型号。",
        "Dryout OFF/ON除tracker外完全同条件，并报告差值。",
        "体积、浓度、孔隙率、Ratio_Dryout和Width均无NaN/Inf与非物理负值。",
        "完成至少两种block大小的收敛性比较，并说明阶梯来源。",
        "容量之外至少用DCR、负极电位、残液/润湿或其他机理证据之一交叉验证。",
        "若提前停止，明确最后有效block、SolverError原文、已完成结果和未覆盖寿命范围。",
    ], numbered=False)
    add_heading(doc, "7.2 常见问题", 2)
    add_body(doc, "问：为什么容量保持率一阶一阶？答：干涸在block之间离散更新；先确认cycles_per_block，再做block-size收敛，不应先用平滑或缩放掩盖。")
    add_body(doc, "问：如何关闭干涸？答：将run_aging_with_dryout的tracker设为None；Notebook建议保留ENABLE_DRYOUT显式开关。")
    add_body(doc, "问：excess_ratio越大是否寿命一定越好？答：它只推迟欠液阈值；若容量由LAM、析锂或其他机理主导，整体寿命未必同步改善。")
    add_body(doc, "问：SolverError是否等于EOL？答：不是。它只说明当前参数/状态下求解失败；必须结合容量、电压、孔隙率、宽度和浓度判断是否为非物理状态。")
    add_body(doc, "问：能否直接使用Electrolyte dry out rate调干涸？答：当前主链不直接读取该字段，应通过SEI耗液、孔隙变化和excess_ratio进行控制，或在明确需求后另行扩展并增加测试。")
    page_break(doc)

    # Closing summary
    add_heading(doc, "基于PyBaMM添加电解液干涸机制的寿命预测模型说明", 1)
    add_heading(doc, "1 编制原因", 2)
    add_body(doc, "本流程说明如何在BatteryProject现有PyBaMM寿命模型中加入电解液干涸准静态耦合，统一物理假设、代码入口、开关、后处理和验收口径，避免把阶梯曲线、历史参数字段或数值崩溃误判为已验证的干涸结论。")
    add_heading(doc, "2 主要内容", 2)
    add_body(doc, "本流程覆盖卷芯孔隙与储液区体积平衡、EC消耗、浓度修正、等效宽度、分块寿命重启、干涸开关对比、可辨识性和项目交付。文中四张图由当前BatteryProject代码及参数生成；容量/DCR寿命结果因缺少指定电芯工况与实验数据而保留待补入口。")
    add_body(doc, "本规范主要起草人：何争。")
    add_body(doc, "本规范为V1版，后续应结合目标电芯实算、残液/拆解证据及update路径集成测试逐步修订。")

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    FINAL.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(FINAL))
    print(FINAL)


if __name__ == "__main__":
    build()
