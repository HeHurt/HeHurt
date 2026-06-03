"""
生成算例文档：《基于PyBaMM添加电解液干涸机制的寿命预测模型》
格式严格参考已有两篇算例的样式和章节结构。
"""
from docx import Document
from docx.shared import Pt, Cm, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import datetime

# ── 用已有文档作为模板以继承自定义样式 ──
TEMPLATE = r"D:\Users\hez\Desktop\hithium\算例\算例\HC_TS_ISC_F413_2025《基于PyBaMM计算电芯不同SOH下能效、功率性能、DCR的方法制定》.docx"
OUTPUT = r"D:\Users\hez\Desktop\hithium\算例\算例\HC_TS_ISC_F413_2025《基于PyBaMM添加电解液干涸机制的寿命预测模型》.docx"

doc = Document(TEMPLATE)

# ── 清空模板全部段落和表格 ──
for p in doc.paragraphs:
    p._element.getparent().remove(p._element)
for t in doc.tables:
    t._element.getparent().remove(t._element)


# ── 辅助函数 ──
def add_para(text, style="Normal", bold=None, size=None, alignment=None, space_after=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if bold is not None:
        run.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    if alignment is not None:
        p.alignment = alignment
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def add_body(text):
    """正文部分样式"""
    try:
        return add_para(text, style="正文部分")
    except KeyError:
        return add_para(text, style="Normal")


def add_caption(text):
    """图表标注样式"""
    try:
        p = add_para(text, style="图表")
    except KeyError:
        p = add_para(text, style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_code_block(lines):
    """模拟代码块：用 Normal 样式，Consolas 字体"""
    for line in lines:
        p = doc.add_paragraph(style="Normal")
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = Pt(13)


def add_table_simple(headers, rows, col_widths=None):
    """添加简单表格"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for par in hdr_cells[i].paragraphs:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in par.runs:
                run.bold = True
                run.font.size = Pt(9)
    for ri, row_data in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row_data):
            cells[ci].text = str(val)
            for par in cells[ci].paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size = Pt(9)
    return table


today = datetime.date.today().strftime("%Y-%m-%d")

# ===================================================================
#  封面
# ===================================================================
add_para("编号：HC_TS_ISC_F413_2025", alignment=WD_ALIGN_PARAGRAPH.LEFT)
add_para("")
add_para("")
add_para("《基于PyBaMM添加电解液干涸机制的寿命预测模型》",
         bold=True, size=18, alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_para("")
add_para("")
add_para(f"编  制：  何 争   日期： {today}", alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_para("")
add_para("校  对：  赖少波  日期： ", alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_para("")
add_para("审  核：  吴长风  日期： ", alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_para("")
add_para("批  准：           日期：", alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_para("")
add_para("厦门海辰储能科技股份有限公司",
         bold=True, size=14, alignment=WD_ALIGN_PARAGRAPH.CENTER)

# ===================================================================
#  编制说明
# ===================================================================
add_para("编制说明", bold=True, size=14, alignment=WD_ALIGN_PARAGRAPH.LEFT)
try:
    bti_style = "Body Text Indent"
    doc.styles[bti_style]
except KeyError:
    bti_style = "Normal"

p = doc.add_paragraph(style=bti_style)
p.add_run("为使本公司仿真业务规范化，参考业界技术要求，结合本公司仿真应用的经验，编制本仿真分析规范。")

p = doc.add_paragraph(style=bti_style)
p.add_run("本规范对本公司仿真计算中心人员在仿真业务开展过程中起到一种指导操作的作用，规范作业流程，固化技术，提高仿真计算效率和精度，对于新进员工展开工作具有关键的指导作用。")

p = doc.add_paragraph(style=bti_style)
p.add_run("本规范对电池寿命预测中电解液干涸(Electrolyte Dry-out)机制的建模方法、开发流程、应用方式及后处理结果作了规定，并在实践中逐步提高完善。")

add_para("本规范主要起草人：何争。")
add_para("本规范的版本记录和版本号变动与修订记录")
add_para("")

# 版本记录表
add_table_simple(
    ["版本号", "制定/修订者", "制定/修订日期", "批准", "日期"],
    [
        ["HC_TS_ISC_F413_2025", "何争", today, "", ""],
    ],
)
add_para("")

# ===================================================================
#  第 1 章  技术情况介绍
# ===================================================================
doc.add_heading("1 技术情况介绍", level=1)

add_body("本规范规定了在PyBaMM寿命预测模型中添加电解液干涸(Electrolyte Dry-out)机制的仿真方法及要求。")
add_body("本规范适用于海辰储能科技股份有限公司所有体系电芯。")

add_para("1.1  仿真工具介绍")
add_body("PyBaMM（Python Battery Mathematical Modeling）是一个用于电池仿真的开源Python库，具备以下特点：")
add_body("模块化架构：支持不同层级的电化学模型，包括单粒子模型（SPM）、伪二维模型（P2D）等。")
add_body("灵活性高：支持自定义电池参数和老化机理，可根据需求调整模型复杂度。")
add_body("计算效率：采用符号计算和稀疏矩阵运算，提高仿真速度，适用于长期寿命预测。")
add_body("可扩展性：可与实验数据结合，实现参数优化和模型校准。")
add_body("在本规范中，采用PyBaMM中的P2D模型，结合自主开发的电解液干涸(Electrolyte Dry-out)模块，用于构建更完整的电池寿命预测模型。")

add_para("1.2  电解液干涸问题背景")
add_body('在锂离子电池循环寿命的后半段，电解液干涸是引起加速容量衰减（即"跳水"或"拐点"）的重要因素之一。其本质是SEI膜的持续生长会不断消耗有机溶剂（如EC），导致卷芯内电解液体积逐步减少。')
add_body("当卷芯内可用电解液不足以完全浸润电极孔隙时，部分电极区域因失去离子传输通道而变为无效面积，表现为等效活性材料损失（LAM）和内阻急剧增大。")
add_body("然而，PyBaMM原生框架中并未内置电解液干涸机制。本规范介绍了基于Li2024论文思路自主开发并集成到BatteryProject框架中的电解液干涸模块，使得寿命预测能覆盖到电解液耗竭阶段。")

# ===================================================================
#  第 2 章  仿真模型介绍
# ===================================================================
doc.add_heading("2 仿真模型介绍", level=1)

add_body("本规范中的基础模型为P2D电化学模型，主要假设如下：")
add_para("1.      仅考虑极片厚度方向的传输过程。")
add_para("2.      颗粒假设为Dv50的球体，忽略颗粒尺寸分布对传输的影响。")
add_para("3.      不考虑热模型，假设电芯在充放电过程中保持恒温。")
add_para("为了更准确预测电芯的容量衰减趋势，模型中考虑了以下主要老化机理：")
add_para("SEI膜增长（ec reaction limited）。")
add_para("SEI导致孔隙率变化。")
add_para("锂沉积（不可逆）。")
add_para("锂沉积导致孔隙率变化。")
add_para("正负极活性颗粒膨胀及破碎。")
add_para("裂纹处SEI再生长。")
add_para("活性材料损失（应力驱动）。")

add_para("2.1  电解液干涸机制原理", bold=True)
add_body("电解液干涸模块的核心物理逻辑如下（实现位于 BatteryProject/src/electrolyte_dryout.py）：")
add_body("（1）SEI生长消耗EC溶剂 → 电解液体积减少")
add_body("（2）SEI/锂析出占据孔隙 → 孔隙体积减少")
add_body("（3）若EC消耗量 > 孔隙减少量 → 需要从储备罐补充电解液")
add_body("（4）若储备罐不足 → 出现干涸（Ratio_Dryout < 1），等效压缩电极宽度")
add_body("（5）若EC消耗量 < 孔隙减少量 → 电解液被挤出到储备罐")
add_body("")

add_para("2.2  核心方程", bold=True)
add_body("干涸模块在每个仿真分块结束后执行一次守恒更新，核心方程如下：")
add_body("")
add_body("① EC溶剂消耗体积：")
add_body("    V_EC,consumed = (ΔLLINegSEI + ΔLLINegSEIcr) × V_m,EC")
add_body("其中 ΔLLINegSEI 为本块仿真中负极SEI的锂损失增量（mol），ΔLLINegSEIcr 为裂纹处SEI的锂损失增量，V_m,EC 为EC的偏摩尔体积（默认 6.684×10⁻⁵ m³/mol）。")
add_body("")
add_body("② 孔隙体积变化：")
add_body("    V_pore,decrease = V_ely,JR,old - V_pore,new")
add_body("    V_pore,new = Σ(ε_i × L_i × L_y × L_z)    (i = 负极、隔膜、正极)")
add_body("其中 ε_i 为各层末态平均孔隙率。")
add_body("")
add_body("③ 电解液净需求量：")
add_body("    V_need = V_EC,consumed - V_pore,decrease")
add_body("")
add_body("④ 分情况处理：")
add_body("    情况A: V_need < 0 → 电解液被挤出，卷芯始终饱和，不干涸。")
add_body("    情况B1: V_need ≥ 0 且储备足 → 完全补液，Ratio_Dryout = 1.0。")
add_body("    情况B2: V_need ≥ 0 且储备不足 → 部分补液，出现干涸：")
add_body("        Ratio_Dryout = V_ely,JR,new / V_pore,new < 1")
add_body("        W_new = Ratio_Dryout × W₀ （电极宽度压缩）")
add_body("")
add_body("⑤ 浓度修正：")
add_body("    电解液中锂离子浓度修正比：Ratio_CeLi = TotLi_new / TotLi_old / Ratio_Dryout")
add_body("    EC溶剂浓度修正比：Ratio_CeEC = TotEC_new / V_JR_new / c_EC_old")
add_body("修正后的浓度通过 apply_dryout_to_initial_conditions() 函数注入到下一段仿真的初始条件中。")

add_para("2.3  模型流程概述", bold=True)
add_body('整个干涸仿真采用"分块仿真 + 块间守恒更新"的方式进行：')
add_body("Phase 1：首圈循环（消除初始SOC偏差）")
add_body("Phase 2：多块老化循环，每块结束后执行干涸更新，流程如下：")
add_body("    1. 刷新老化加速参数")
add_body("    2. tracker.update(sol, params) → 计算干涸守恒更新")
add_body("    3. apply_dryout_to_initial_conditions(model, sol, params) → 修正初始条件")
add_body("    4. 用更新后的 params 启动下一块仿真")

# ===================================================================
#  第 3 章  仿真工况及假设条件
# ===================================================================
doc.add_heading("3 仿真工况及假设条件", level=1)

add_body("工作温度区间：放电-20~60℃、充电0℃~60℃（通用情况）")
add_body("仿真工况：恒功率循环，充放间隔10分钟静置")
add_body("循环工况示例：")
add_body("    充电：940W充电至3.65V")
add_body("    静置：10分钟")
add_body("    放电：940W放电至2.5V")
add_body("    静置：10分钟")
add_body("干涸机制假设条件：")
add_body("    1. EC消耗与SEI生长成正比，简化为 EC:Li:SEI = 1:1 的摩尔关系")
add_body("    2. 电解液为不可压缩流体，储备罐与卷芯间自由流动")
add_body("    3. 干涸引起的活性面积减少等效为电极宽度的线性压缩")
add_body("    4. 块间更新时假设各域电解液浓度瞬间均匀混合")

# ===================================================================
#  第 4 章  设计/校对输入
# ===================================================================
doc.add_heading("4 设计/校对输入", level=1)

add_body("寿命仿真分析（含干涸机制）需求的输入数据如表4-1：")
add_caption("表4-1 设计/校对输入")

add_table_simple(
    ["序号", "输入项目", "参数/数模", "归属", "备注"],
    [
        ["1", "电池设计参数", "参数", "电化学模型", "极片厚度、孔隙率、电极宽度/高度等"],
        ["2", "正负极OCP/OCV曲线", "参数", "电化学模型", "充放电方向各一条"],
        ["3", "电解质参数", "参数", "电化学模型", "扩散系数、电导率等"],
        ["4", "初始电解液过量比", "参数", "干涸模型", "excess_ratio, 通常1.0~1.5"],
        ["5", "EC偏摩尔体积", "参数", "干涸模型", "默认6.684×10⁻⁵ m³/mol"],
        ["6", "老化动力学参数", "参数", "老化模型", "SEI速率、EC扩散率、镀锂速率等"],
        ["7", "实测容量保持率曲线", "数据", "校准", "用于老化参数对标"],
    ],
)

# ===================================================================
#  第 5 章  流程及方法
# ===================================================================
doc.add_heading("5 流程及方法", level=1)

add_body("电解液干涸寿命预测方法制定流程如下：")
add_body("步骤1：建立P2D电化学基础模型，完成BOL性能校准")
add_body("步骤2：开启老化机制，完成短期老化曲线对标")
add_body("步骤3：初始化DryoutTracker，设定电解液过量比")
add_body("步骤4：首圈仿真（消除初始SOC偏差）")
add_body("步骤5：分块循环 + 块间干涸更新")
add_body("步骤6：后处理与结果分析")

doc.add_heading("5.1 软硬件要求", level=2)
doc.add_heading("5.1.1  软件要求", level=3)
add_body("前处理：Visual Studio Code")
add_body("求解器：PyBaMM内置Idaklu")
add_body("后处理：Jupyter Notebook、Origin")
add_body("干涸模块：BatteryProject/src/electrolyte_dryout.py")
doc.add_heading("5.1.2  硬件要求", level=3)
add_body("对于前后处理，本地工作站是唯一选择；对于求解的过程，可以在工作站或者高性能服务器上进行。")

doc.add_heading("5.2 仿真参数设置流程", level=2)
doc.add_heading("5.2.1 电化学模型建立", level=3)
add_body("1、仿真需求提出后，需项目组提供电池设计参数表，填写对应体系各项参数。")
add_body("2、确认设计参数无误后，将参数表中各项参数对应填写至PyBaMM中，建立起对应电芯体系的1D电化学模型。")
add_body("3、DFN模型需开启以下老化选项（干涸机制依赖SEI和孔隙率变化数据）：")
add_code_block([
    'model = pybamm.lithium_ion.DFN({',
    '    "SEI": "ec reaction limited",',
    '    "SEI porosity change": "true",',
    '    "lithium plating": "irreversible",',
    '    "lithium plating porosity change": "true",',
    '    "particle mechanics": ("swelling and cracking", "swelling only"),',
    '    "SEI on cracks": "true",',
    '    "loss of active material": "stress-driven",',
    '    "calculate discharge energy": "true",',
    '    "contact resistance": "true",',
    '    "open-circuit potential": ("single", "current sigmoid"),',
    '})',
])
add_caption("代码5-1 DFN模型老化可选项设置")
add_body("")
add_body("4、设置求解器及网格划分：")
add_code_block([
    'solver = pybamm.IDAKLUSolver()',
    'var_pts = {"x_n": 5, "x_s": 5, "x_p": 5, "r_n": 20, "r_p": 20}',
])
add_caption("代码5-2 求解器及网格设置")

doc.add_heading("5.2.2 干涸模块集成", level=3)
add_body("完成基础电化学模型校准和老化参数对标后，即可集成电解液干涸机制。以下为完整的Notebook操作流程：")
add_body("")

add_body("步骤1：导入干涸模块")
add_code_block([
    '# 基础环境与快速导入',
    '%load_ext autoreload',
    '%autoreload 2',
    'import sys',
    'from pathlib import Path',
    'import pybamm',
    'import numpy as np',
    '',
    'PROJECT_ROOT = Path(r"D:\\Users\\hez\\Desktop\\hithium\\BatteryProject")',
    'if str(PROJECT_ROOT) not in sys.path:',
    '    sys.path.append(str(PROJECT_ROOT))',
    '',
    'from src.easy_imports import *',
    '# 以上会自动导入 DryoutTracker, apply_dryout_to_initial_conditions,',
    '# plot_dryout, run_aging_with_dryout 等函数',
])
add_caption("代码5-3 导入干涸模块")

add_body("步骤2：设置干涸关键参数并初始化追踪器")
add_code_block([
    '# ============ 干涸关键参数 ============',
    'excess_ratio = 1.2          # 初始电解液过量比 (>1 有储备余量)',
    'n_blocks = 50               # 分块数',
    'cycles_per_block = 2        # 每块循环圈数',
    't_factor_aging = 50         # 老化加速因子',
    '# ======================================',
    '',
    'params = pybamm.ParameterValues("OKane2022")',
    'params.update(get_hithium_params(1, temperature=temperature),',
    '              check_already_exists=False)',
    '',
    '# 初始化干涸追踪器',
    'tracker = DryoutTracker(params, excess_ratio=excess_ratio)',
])
add_caption("代码5-4 干涸关键参数与追踪器初始化")

add_body("DryoutTracker构造函数会自动执行以下操作：")
add_body("    a. 从params中读取极片厚度、电极宽高、各层孔隙率")
add_body("    b. 计算卷芯初始孔隙体积 V_pore = Σ(L_i × ε_i × L_y × L_z)")
add_body("    c. 按 excess_ratio 计算总电解液量 V_tot = V_pore × excess_ratio")
add_body("    d. 向params注入追踪所需的键值（如当前卷芯内电解液体积、储备罐浓度等）")

add_body("")
add_body("步骤3：首圈仿真（消除初始SOC偏差）")
add_code_block([
    'exp_init = pybamm.Experiment([',
    '    (',
    '        f"Charge at {rate*nominal_current*nominal_voltage:.0f}W '
    'until 3.65 V (0.5 minute period)",',
    '        "Rest for 10 minute (0.5 minute period)",',
    '        f"Discharge at {rate*nominal_current*nominal_voltage:.0f}W '
    'until 2.5 V (0.5 minute period)",',
    '        "Rest for 10 minute (0.5 minute period)",',
    '    )',
    '], temperature=temperature)',
    '',
    'sim_init = pybamm.Simulation(model, parameter_values=params,',
    '                              experiment=exp_init,',
    '                              var_pts=var_pts, solver=solver)',
    'sol = sim_init.solve(showprogress=True)',
])
add_caption("代码5-5 首圈仿真")

add_body("")
add_body("步骤4：分块老化循环（核心干涸更新逻辑）")
add_code_block([
    'dry_sol_list = []',
    '',
    'for i_block in range(n_blocks):',
    '    print(f"\\n=== Block {i_block+1}/{n_blocks} ===")',
    '',
    '    # 刷新老化参数（加速因子）',
    '    params.update(get_hithium_params(t_factor_aging,',
    '                  temperature=temperature),',
    '                  check_already_exists=False)',
    '',
    '    # 干涸更新: 计算 EC 消耗、浓度修正',
    '    tracker.update(sol, params)',
    '',
    '    # 将干涸修正应用到模型初始条件',
    '    model_run = apply_dryout_to_initial_conditions(',
    '        model, sol, params)',
    '',
    '    # 本块实验',
    '    exp_block = pybamm.Experiment([',
    '        (',
    '            f"Charge at {rate*nominal_current*nominal_voltage:.0f}W '
    'until 3.65 V (0.5 minute period)",',
    '            "Rest for 10 minute (0.5 minute period)",',
    '            f"Discharge at {rate*nominal_current*nominal_voltage:.0f}W '
    'until 2.5 V (0.5 minute period)",',
    '            "Rest for 10 minute (0.5 minute period)",',
    '        )',
    '    ] * cycles_per_block, temperature=temperature)',
    '',
    '    sim_block = pybamm.Simulation(model_run,',
    '        experiment=exp_block,',
    '        parameter_values=params,',
    '        solver=solver, var_pts=var_pts)',
    '    sol = sim_block.solve(starting_solution=sol,',
    '        showprogress=True, calc_esoh=False)',
    '    dry_sol_list.append(sol)',
    '',
    '# 最后一块结束后再更新一次追踪',
    'tracker.update(sol, params)',
    'tracker.summary()',
])
add_caption("代码5-6 分块老化循环（带干涸更新）")
add_body("每个block之间，tracker.update()内部调用_cal_new_con_update()函数计算EC消耗量、孔隙变化量、补液/挤出量，并就地更新params中的电解液体积、浓度比例和电极宽度。")

doc.add_heading("5.2.3 关键参数说明与调参指南", level=3)
add_body("以下表格汇总了影响干涸速度的关键参数及其调参方向：")
add_caption("表5-1 干涸相关关键参数一览")

add_table_simple(
    ["参数名", "默认值", "作用", "加速干涸方向"],
    [
        ["excess_ratio", "1.2", "初始电解液过量比，决定储备罐容量", "↓ 减小"],
        ["t_factor_aging", "50", "老化加速因子，放大SEI/镀锂/裂纹等", "↑ 增大"],
        ["temperature", "298.15 K", "温度，影响副反应Arrhenius速率", "↑ 升高"],
        ["SEI kinetic rate constant", "4.8e-14 m/s", "SEI生成动力学常数", "↑ 增大"],
        ["EC diffusivity", "3.5e-22 m²/s", "EC在SEI中的扩散系数", "↑ 增大"],
        ["Lithium plating kinetic rate", "3e-5 m/s", "锂沉积速率常数", "↑ 增大"],
        ["Negative electrode cracking rate", "~2.2e-21", "颗粒裂纹扩展速率", "↑ 增大"],
        ["EC partial molar volume", "6.684e-5 m³/mol", "EC偏摩尔体积", "↑ 增大"],
        ["cycles_per_block", "2", "每块循环数，影响反馈频率", "↓ 减小（更频繁反馈）"],
        ["n_blocks", "50", "分块数量，影响总循环数", "↑ 增大"],
    ],
)

add_body("")
add_body("调参注意事项：")
add_body("1. excess_ratio 是最直接、最安全的独立旋钮。建议从 1.2 开始，分别尝试 1.1 和 1.05 来观察干涸曲线变化。")
add_body('2. t_factor_aging 会同步放大所有老化机制（SEI + 镀锂 + 裂纹 + LAM），适合"整体加速"场景，但不能单独控制干涸。')
add_body("3. 温度是间接杠杆：升温会通过Arrhenius方程加速所有副反应动力学。")
add_body("4. 需要独立调控干涸速度时，优先调 excess_ratio；如需更精细控制可修改 EC partial molar volume 或向 _cal_new_con_update 中添加乘数系数。")
add_body("5. 注意：config.py 中的 'Electrolyte dry out rate [m3.s-1]' 参数当前未被干涸主流程读取，修改该值不会影响干涸速度。")

# ===================================================================
#  第 6 章  后处理及功能展示
# ===================================================================
doc.add_heading("6 后处理及功能展示", level=1)

add_body("在干涸仿真完成后，可进行以下后处理分析：")

doc.add_heading("6.1 干涸演化可视化", level=2)
add_body("DryoutTracker内置了 plot() 方法，可一键输出 2×3 的综合面板图，包含以下6个子图：")
add_body("    子图1：电解液体积（Total / Jelly Roll / Pore Volume）")
add_body("    子图2：干涸比例 Ratio_Dryout 随循环变化")
add_body("    子图3：电极宽度 Width 随循环变化")
add_body("    子图4：EC消耗量、补液量")
add_body("    子图5：浓度比例（CeEC ratio / CeLi ratio）")
add_body("    子图6：储备罐浓度（Li⁺ / EC）")
add_body("")
add_code_block([
    'tracker.plot()',
])
add_caption("代码6-1 干涸演化可视化")
add_body("典型输出结果：当excess_ratio=1.2时，在前30~40个block中Ratio_Dryout保持1.0（储备充足），在储备耗尽后Ratio_Dryout开始下降，电极宽度同步收缩。")

doc.add_heading("6.2 容量保持率曲线", level=2)
add_body("使用BatteryPlotter结合process_sol_list_with_custom_extractor，可以将干涸仿真结果叠加到容量保持率图中：")
add_code_block([
    'my_plotter = BatteryPlotter()',
    '',
    '# 注入干涸仿真的容量数据',
    'process_sol_list_with_custom_extractor(',
    '    my_plotter,',
    '    dry_sol_list,',
    '    [f"Dryout {rate}P"] * len(dry_sol_list),',
    '    t_factor=t_factor_aging,',
    ')',
    '',
    'my_plotter.plot(target_sim="all", mode="capacity")',
])
add_caption("代码6-2 容量保持率绘图")
add_body("加入干涸机制后的典型表现：早期容量保持率曲线与无干涸版本基本重合，在储备耗尽后出现加速衰减拐点。")

doc.add_heading("6.3 干涸等效LAM分析", level=2)
add_body("干涸引起的电极宽度压缩可等效换算为LAM百分比，计算方式为：")
add_body("    LAM_dry(%) = 100 - (W_current / W_initial) × 100")
add_code_block([
    'lam_dry_pct = tracker.get_lam_dryout_pct()',
    'steps = np.arange(len(lam_dry_pct))',
    '',
    'fig, ax = plt.subplots(figsize=(8, 4))',
    'ax.plot(steps, lam_dry_pct, "o-", color="tab:red",',
    '        label="LAM from dryout")',
    'ax.set_xlabel("Step")',
    'ax.set_ylabel("LAM (%)")',
    'ax.set_title("LAM Contribution from Electrolyte Dry-out")',
    'ax.legend()',
    'ax.grid(True, alpha=0.3)',
    'plt.tight_layout()',
    'plt.show()',
])
add_caption("代码6-3 干涸等效LAM分析")

doc.add_heading("6.4 综合分析图", level=2)
add_body("调用 plot_analysis(sol, t_factor_aging) 可输出包含寿命、各老化分量占比、内部变量演化等的综合分析图。")
add_code_block([
    'plot_analysis(sol, t_factor_aging)',
])
add_caption("代码6-4 综合分析图")

doc.add_heading("6.5 批量参数扫描对比", level=2)
add_body("为快速评估不同参数设置对干涸的影响，干涸Notebook中提供了批量扫描功能。配置多组案例后，一次运行即可输出对比图：")
add_code_block([
    '# 配置扫描案例',
    'batch_cases = [',
    '    {"name": "ER=1.20", "excess_ratio": 1.20},',
    '    {"name": "ER=1.10", "excess_ratio": 1.10},',
    '    {"name": "ER=1.05", "excess_ratio": 1.05},',
    ']',
    '',
    '# 执行扫描',
    'scan_results = []',
    'for case in batch_cases:',
    '    merged = {**scan_base_case, **case}',
    '    scan_results.append(',
    '        run_dryout_scan_case(merged, showprogress=True))',
    '',
    '# 输出三联对比图',
    '# - Dryout Ratio 对比',
    '# - Dryout-induced LAM 对比',
    '# - Capacity Retention 对比',
])
add_caption("代码6-5 批量参数扫描")
add_body("典型扫描结果：excess_ratio从1.20降至1.05时，干涸拐点提前约200~300等效圈，LAM从~2%增至~12%，容量保持率曲线出现更明显的尾部跳水。")

doc.add_heading("6.6 一键封装接口", level=2)
add_body("对于不需要逐块控制逻辑的场景，可使用封装好的 run_aging_with_dryout() 函数一键完成整个干涸老化仿真：")
add_code_block([
    'from src.electrolyte_dryout import run_aging_with_dryout',
    '',
    'sol_list = run_aging_with_dryout(',
    '    model=model,',
    '    params=params,',
    '    experiment=make_experiment,  # callable(cycles)',
    '    solver=solver,',
    '    var_pts=var_pts,',
    '    tracker=tracker,',
    '    n_blocks=50,',
    '    cycles_per_block=2,',
    '    t_factor=50,',
    '    temperature=298.15,',
    '    get_hithium_params=get_hithium_params,',
    ')',
])
add_caption("代码6-6 一键干涸老化仿真接口")

# ===================================================================
#  第 7 章  项目组反馈答疑
# ===================================================================
doc.add_heading("7 项目组反馈答疑", level=1)

add_body("当所有结果汇总完成后，名称命名为（Hithium_TI_xxxx Cell XXXX_日期_Vx），经部门内确认校准后，交付给项目组验收。")
add_body("向项目组交付结果时需要一并将仿真假设条件、简化之处告知项目组，重点说明以下事项：")
add_body("1. 干涸机制属于块间准静态更新，非求解器内连续耦合，精度随 cycles_per_block 增大而降低。")
add_body("2. 干涸主要影响寿命后半段的容量衰减斜率和拐点位置，对BOL性能无影响。")
add_body("3. excess_ratio 需结合电芯实际注液量和设计孔隙率估算，不可随意猜测。")
add_body("4. 当前模型将干涸等效为电极宽度线性压缩，未考虑局部干涸引起的非均匀老化。")
add_body("以上流程为基于PyBaMM添加电解液干涸机制的寿命预测模型通用操作规范，当项目组提出超出一般精度或者常规仿真的需求时，需积极与项目组沟通。")

# ===================================================================
#  保存
# ===================================================================
doc.save(OUTPUT)
print(f"文档已保存: {OUTPUT}")
