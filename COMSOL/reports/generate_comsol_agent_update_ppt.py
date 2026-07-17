from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Cm, Pt


OUT = Path(__file__).with_name("COMSOL_Agent_Hybrid工具链更新汇报_20260623.pptx")

NAVY = RGBColor(11, 36, 71)
RED = RGBColor(200, 16, 46)
BLUE = RGBColor(46, 108, 196)
AMBER = RGBColor(245, 166, 35)
GRAY = RGBColor(107, 114, 128)
LIGHT = RGBColor(244, 246, 249)
TEXT = RGBColor(31, 41, 55)
WHITE = RGBColor(255, 255, 255)
PALE_BLUE = RGBColor(232, 240, 253)
PALE_RED = RGBColor(253, 235, 239)
PALE_GREEN = RGBColor(232, 246, 239)


def rgb(hexstr: str) -> RGBColor:
    hexstr = hexstr.strip("#")
    return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))


def set_text_style(paragraph, size=18, color=TEXT, bold=False, font="Microsoft YaHei"):
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold


def text_box(slide, x, y, w, h, text="", size=18, color=TEXT, bold=False, align=None):
    shape = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    set_text_style(p, size=size, color=color, bold=bold)
    return shape


def add_title(slide, title, subtitle=None, dark=False):
    color = WHITE if dark else NAVY
    text_box(slide, 1.0, 0.55, 24.0, 0.9, title, size=25, color=color, bold=True)
    if subtitle:
        text_box(slide, 1.05, 1.42, 23.0, 0.55, subtitle, size=11.5, color=(WHITE if dark else GRAY))
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.0), Cm(2.16), Cm(3.0), Cm(0.08)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = RED
    slide.shapes[-1].line.fill.background()


def add_footer(slide, n):
    text_box(slide, 1.0, 18.35, 12.0, 0.35, "Hithium COMSOL Agent Workflow", size=7.5, color=GRAY)
    text_box(slide, 24.5, 18.35, 1.5, 0.35, f"{n:02d}", size=8, color=GRAY, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body, fill=LIGHT, accent=BLUE, title_size=13, body_size=10.5):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = rgb("D9DEE7")
    s.line.width = Pt(0.8)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(0.13), Cm(h)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = accent
    slide.shapes[-1].line.fill.background()
    text_box(slide, x + 0.35, y + 0.25, w - 0.6, 0.45, title, size=title_size, color=accent, bold=True)
    text_box(slide, x + 0.35, y + 0.85, w - 0.6, h - 1.05, body, size=body_size, color=TEXT)
    return s


def bullet_box(slide, x, y, w, h, items, size=12, color=TEXT, fill=None, title=None):
    if fill:
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
        s.fill.solid()
        s.fill.fore_color.rgb = fill
        s.line.color.rgb = rgb("E4E7EC")
    tx = slide.shapes.add_textbox(Cm(x + 0.25), Cm(y + 0.2), Cm(w - 0.5), Cm(h - 0.3))
    tf = tx.text_frame
    tf.clear()
    tf.word_wrap = True
    if title:
        p = tf.paragraphs[0]
        p.text = title
        set_text_style(p, size=size + 1, color=NAVY, bold=True)
    else:
        p = tf.paragraphs[0]
        p.text = items[0]
        p.level = 0
        set_text_style(p, size=size, color=color)
        items = items[1:]
    for item in items:
        p = tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(3)
        p.margin_left = Cm(0.28)
        set_text_style(p, size=size, color=color)
    return tx


def arrow(slide, x1, y1, x2, y2, color=GRAY):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    line.line.color.rgb = color
    line.line.width = Pt(1.4)
    line.line.end_arrowhead = True
    return line


def pill(slide, x, y, w, text, fill, color=WHITE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(0.72))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    text_box(slide, x + 0.1, y + 0.12, w - 0.2, 0.45, text, size=9.5, color=color, bold=True, align=PP_ALIGN.CENTER)
    return s


def make_deck():
    prs = Presentation()
    prs.slide_width = Cm(26.67)
    prs.slide_height = Cm(15.0)
    blank = prs.slide_layouts[6]

    # 1 Cover
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(13.2), Cm(26.67), Cm(1.8)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = rgb("071A33")
    slide.shapes[-1].line.fill.background()
    text_box(slide, 1.2, 1.6, 22.5, 1.2, "COMSOL Agent Hybrid 工具链更新汇报", size=27, color=WHITE, bold=True)
    text_box(slide, 1.25, 3.0, 22.5, 0.8, "从 Java/batch 主链到 sim-plugin-comsol + MCP 观察层", size=16, color=rgb("D9E2F3"))
    pill(slide, 1.25, 4.3, 4.5, "Skill v3.1", RED)
    pill(slide, 6.1, 4.3, 5.8, "sim-cli-core 0.3.7", BLUE)
    pill(slide, 12.3, 4.3, 6.5, "sim-plugin-comsol 0.1.14", AMBER, color=NAVY)
    text_box(slide, 1.25, 12.7, 17, 0.45, "2026-06-23  |  Hithium 电池仿真项目", size=10.5, color=rgb("D9E2F3"))
    text_box(slide, 21.4, 0.55, 4.2, 0.45, "HITHIUM", size=15, color=WHITE, bold=True, align=PP_ALIGN.RIGHT)

    # 2 Why update
    slide = prs.slides.add_slide(blank)
    add_title(slide, "为什么要更新", "原 Java/batch 主链能交付，但缺少对现有 .mph、运行态和人工协作的观察能力")
    card(slide, 1.0, 2.8, 7.4, 4.5, "已具备", "AI 生成/修改 COMSOL Java\ncomsolcompile 编译\ncomsolbatch 求解\n输出 .mph / metrics / report", fill=PALE_GREEN, accent=BLUE)
    card(slide, 9.6, 2.8, 7.4, 4.5, "主要缺口", "已有 .mph 是黑盒\nDLP/TSD 加密影响普通进程读取\n边界、dataset、result 节点不可见\n调试只能依赖日志", fill=PALE_RED, accent=RED)
    card(slide, 18.2, 2.8, 7.4, 4.5, "更新目标", "让 Agent 不只会“写模型”\n还要能“看模型、查状态、协同桌面、快速后处理”\n但不牺牲 Java artifact 可复现性", fill=PALE_BLUE, accent=BLUE)
    text_box(slide, 1.1, 8.5, 24.5, 1.1, "核心判断：Java 是可复现交付底座；sim/MCP 是观察与操作增强层，不是替代 Java 主链。", size=16, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 2)

    # 3 Architecture
    slide = prs.slides.add_slide(blank)
    add_title(slide, "更新后的分层原理", "三层各司其职，避免把所有能力塞进一种接口")
    labels = [
        ("需求/验收", "用户定义目标、参数、边界条件、输出指标", NAVY),
        ("AI Agent", "理解需求、规划建模、解析报错、决定下一轮迭代", RED),
        ("Java 主链", ".java / comsolcompile / comsolbatch / .mph", BLUE),
        ("sim Runtime", "comsolmphserver / JPype / shared Desktop", AMBER),
        ("MCP 查询层", "参数、边界、dataset、结果、导出", rgb("2A9D8F")),
        ("报告交付", "metrics / plots / debug history / docx or pptx", NAVY),
    ]
    xs = [1.0, 5.1, 9.2, 13.3, 17.4, 21.5]
    for i, (title, desc, color) in enumerate(labels):
        card(slide, xs[i], 3.5, 3.5, 4.3, title, desc, fill=LIGHT, accent=color, title_size=12, body_size=8.8)
        if i < len(labels) - 1:
            arrow(slide, xs[i] + 3.5, 5.65, xs[i + 1], 5.65, color=GRAY)
    bullet_box(
        slide,
        1.0,
        9.0,
        24.7,
        2.2,
        [
            "Java/batch：默认主链，负责可复现模型与交付物。",
            "sim-plugin-comsol：负责 COMSOL 运行时、常驻会话、shared Desktop、加密 .mph 读取。",
            "MCP：负责结构化查询与临时后处理，不能跳过 Java artifact。",
        ],
        size=11.5,
        fill=rgb("F7F8FA"),
    )
    add_footer(slide, 3)

    # 4 Skill updates
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Skill 更新了什么", "comsol-java-battery-modeling 从 v3.0 升级到 v3.1")
    card(slide, 1.0, 2.4, 7.8, 8.8, "1. 增加 Hybrid 工具层", "保留 Snippet / Autonomous 两种交付模式。\n新增 sim-cli、sim-plugin-comsol、COMSOL MCP 的启用边界。\n原则：先判断缺的是生成、运行时还是观察能力。", fill=PALE_BLUE, accent=BLUE)
    card(slide, 9.45, 2.4, 7.8, 8.8, "2. 增加路由决策表", "从零建模：Java/batch。\n已有 .mph：inspect / MCP。\n人机协作：shared Desktop。\n临时结果读取：MCP。\n批量重跑：comsolbatch。", fill=PALE_GREEN, accent=rgb("2A9D8F"))
    card(slide, 17.9, 2.4, 7.8, 8.8, "3. 加入红线和 STEP 补充", "STEP A 增加 Hybrid 前置确认。\nSTEP E 增加 MCP 后处理边界。\n红线：不能为了 MCP 跳过 .java/.mph/log/metrics；不能暴露 sim serve。", fill=PALE_RED, accent=RED)
    add_footer(slide, 4)

    # 5 Environment updates
    slide = prs.slides.add_slide(blank)
    add_title(slide, "环境与脚本更新", "把验证链路固化，避免每次 --with 冷启动探索")
    card(slide, 1.0, 2.35, 11.8, 3.2, "根目录 pyproject.toml", "新增轻量 agent-tools 环境：\nsim-cli-core==0.3.7\nsim-plugin-comsol==0.1.14", fill=LIGHT, accent=BLUE)
    card(slide, 13.8, 2.35, 11.8, 3.2, "tools/comsol_mph_probe.py", "固定 .mph 读取入口：\n复用 default session；无 session 时自动 no_gui connect；输出文件头、TSD 判断、模型树摘要。", fill=LIGHT, accent=RED)
    card(slide, 1.0, 6.3, 11.8, 3.2, "清理环境噪声", "删除系统 Python 旧 editable 残留：\nBatterModel.egg-link\n从 easy-install.pth 移除 BatteryModelPack_v1.4 路径。", fill=LIGHT, accent=AMBER)
    card(slide, 13.8, 6.3, 11.8, 3.2, "使用注意", "不要并行跑多个 uv run sim。\nCOMSOL 冷启动约分钟级；session 保持后读取小 .mph 约 10-20 秒。", fill=LIGHT, accent=rgb("2A9D8F"))
    text_box(slide, 1.0, 10.75, 24.0, 0.75, "新增文件：pyproject.toml、tools/comsol_mph_probe.py；skill 文件：skills/comsol-java-battery-modeling/SKILL.md。", size=11, color=GRAY)
    add_footer(slide, 5)

    # 6 Validation results
    slide = prs.slides.add_slide(blank)
    add_title(slide, "已验证的新能力", "本次验证不是理论设计，已用本机 COMSOL 6.4 和 TSD 样本跑通")
    headers = ["能力", "验证方式", "结果"]
    rows = [
        ["工具识别", "uv run sim check comsol", "识别 COMSOL 6.4 / comsol_64_mph_1"],
        ["TSD .mph 读取", "ModelUtil.load 加载 张倩.mph / 复现粒径分布.mph", "成功；COMSOL session 可读加密文件"],
        ["模型树探针", "读取 张倩.mph 内部节点", "comp1 / geom1 / es / mesh1 / std1 / pg1 pg2"],
        ["shared Desktop", "visual_mode=shared-desktop + session.health", "model_builder_live=true；live_model_binding.ok=true"],
        ["桌面协作建模", "Agent 执行 block_with_hole 几何步骤", "Desktop 绑定 Model1；读回 comp1 / geom1"],
    ]
    table = slide.shapes.add_table(len(rows) + 1, 3, Cm(1.0), Cm(2.4), Cm(24.7), Cm(6.6)).table
    table.columns[0].width = Cm(4.9)
    table.columns[1].width = Cm(10.8)
    table.columns[2].width = Cm(9.0)
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        set_text_style(p, size=10.5, color=WHITE, bold=True)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if i % 2 else RGBColor(255, 255, 255)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            set_text_style(p, size=9.2, color=TEXT)
    card(slide, 1.0, 10.0, 7.7, 2.1, "速度基线", "plugin list 约 4.6s\nsim ps 约 9.1s\n显式 --session probe 约 11.3s", fill=PALE_BLUE, accent=BLUE, title_size=11, body_size=9.5)
    card(slide, 9.5, 10.0, 7.7, 2.1, "关键结论", "加密文件不是 Agent 直接解密；是 COMSOL/mphserver 进程链具备明文读取权限。", fill=PALE_GREEN, accent=rgb("2A9D8F"), title_size=11, body_size=9.5)
    card(slide, 18.0, 10.0, 7.7, 2.1, "边界", "普通手开 COMSOL Desktop 不是 shared session；需由 sim connect 启动绑定桌面。", fill=PALE_RED, accent=RED, title_size=11, body_size=9.5)
    add_footer(slide, 6)

    # 7 How to use
    slide = prs.slides.add_slide(blank)
    add_title(slide, "后续如何使用", "按任务场景选择最小工具集合")
    card(slide, 1.0, 2.25, 7.7, 7.4, "后台快速读 .mph", "适用：查看模型结构、读取 TSD 加密样本。\n\n命令：\nuv run python tools\\comsol_mph_probe.py --session <id> \"D:\\path\\model.mph\"\n\n无 session 时可省略 --session，脚本会自动 no_gui connect。", fill=LIGHT, accent=BLUE, title_size=12.5, body_size=9.2)
    card(slide, 9.5, 2.25, 7.7, 7.4, "边看边建模", "适用：工程师要观察 Model Builder。\n\n命令：\nuv run sim connect --solver comsol --ui-mode gui --driver-option visual_mode=shared-desktop\nuv run sim inspect session.health\n\n确认 live_model_binding.ok=true 后再 exec。", fill=LIGHT, accent=RED, title_size=12.5, body_size=9.2)
    card(slide, 18.0, 2.25, 7.7, 7.4, "可复现交付", "适用：正式建模、批量扫描、报告。\n\n主链不变：\nAI 写/改 .java\ncomsolcompile 编译\ncomsolbatch 求解\n输出 .mph / metrics / debug_history / report\n\nMCP/sim 只做增强。", fill=LIGHT, accent=rgb("2A9D8F"), title_size=12.5, body_size=9.2)
    text_box(slide, 1.1, 10.6, 24.2, 0.9, "操作原则：冷启动只做一次；session 保持期间复用同一个 id；不要并行抢 uv/sim 锁；固定验收指标写回 Java/export。", size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 7)

    # 8 conclusion
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    add_title(slide, "结论与下一步", "本次更新把 COMSOL Agent 从“能生成代码”推进到“能观察、协作和复用会话”", dark=True)
    card(slide, 1.1, 3.0, 7.6, 4.7, "已经完成", "Skill v3.1 路由规则\n项目级 sim 固定依赖\n.mph probe 固定入口\nTSD 加密 .mph 读取验证\nshared Desktop 绑定验证", fill=rgb("123354"), accent=AMBER, title_size=13, body_size=10.5)
    card(slide, 9.55, 3.0, 7.6, 4.7, "使用边界", "Java/batch 仍是正式交付主链\nsim 负责运行时与 Desktop\nMCP 负责结构化查询和临时后处理\n三者不机械串联", fill=rgb("123354"), accent=RED, title_size=13, body_size=10.5)
    card(slide, 18.0, 3.0, 7.6, 4.7, "建议下一步", "把 probe 输出接入报告模板\n补充常用 .mph 节点摘要格式\n建立 shared Desktop 分步建模 SOP\n再接 MCP 结果读取 smoke test", fill=rgb("123354"), accent=BLUE, title_size=13, body_size=10.5)
    text_box(slide, 1.1, 10.2, 24.5, 1.2, "目标状态：用户只定义需求、参数、边界和验收指标；Agent 自动建模、求解、观察、迭代、后处理并输出报告。", size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, 1.2, 13.1, 22, 0.45, "COMSOL Agent Hybrid Workflow Update", size=10, color=rgb("D9E2F3"))
    text_box(slide, 24.2, 13.1, 1.3, 0.45, "08", size=8, color=rgb("D9E2F3"), align=PP_ALIGN.RIGHT)

    prs.save(OUT)


if __name__ == "__main__":
    make_deck()
    print(OUT)
