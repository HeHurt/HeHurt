"""Lifecycle experiment data preparation for aging model comparison."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

import numpy as np
import openpyxl
import pandas as pd


DAT_COLUMNS = [
    "curve_type",
    "cell_id",
    "label",
    "temperature_c",
    "rate_p",
    "cycle",
    "cycle_raw",
    "charge_capacity_ah",
    "discharge_capacity_ah",
    "charge_energy_wh",
    "discharge_energy_wh",
    "capacity_retention_nominal",
    "energy_retention_nominal",
    "energy_efficiency",
    "capacity_retention_bol",
    "q_bol_ah",
    "bol_cycle_raw",
    "n_cells",
]


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and np.isfinite(float(value))


def _temperature_text(temperature_c: float) -> str:
    return f"{temperature_c:g}°C"


def _rate_text(rate_p: float) -> str:
    return f"{rate_p:g}P"


def _condition_label(temperature_c: float, rate_p: float, cell_id: str | None = None) -> str:
    base = f"{_temperature_text(temperature_c)} {_rate_text(rate_p)}"
    return f"{base} {cell_id}" if cell_id else base


def _parse_cell_header(header: str) -> tuple[str, float]:
    match = re.search(r"(Cell\s*\d+).*?(-?\d+(?:\.\d+)?)\s*[°℃]?\s*[cC℃]", header)
    if not match:
        raise ValueError(f"Cannot parse cell header: {header!r}")
    cell_id = re.sub(r"\s+", " ", match.group(1)).strip()
    temperature_c = float(match.group(2))
    return cell_id, temperature_c


def _read_sheet_rows(workbook_path: str | Path, sheet_name: str) -> list[list[object]]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet {sheet_name!r} not found. Available sheets: {workbook.sheetnames}")
        worksheet = workbook[sheet_name]
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _cell_block_starts(rows: list[list[object]]) -> list[tuple[int, str]]:
    if len(rows) < 4:
        raise ValueError("Cycle-life sheet must contain at least 4 rows")
    starts = []
    for col_idx, value in enumerate(rows[0]):
        if isinstance(value, str) and value.strip().lower().startswith("cell"):
            starts.append((col_idx, value.strip()))
    if not starts:
        raise ValueError("No measured cell blocks found in the first row")
    return starts


def _parse_measured_blocks(
    rows: list[list[object]],
    *,
    nominal_capacity_ah: float,
    nominal_energy_wh: float,
    rate_p: float,
) -> pd.DataFrame:
    records: list[dict[str, object]] = []
    for start_col, header in _cell_block_starts(rows):
        cell_id, temperature_c = _parse_cell_header(header)
        raw_rows = []
        for row in rows[3:]:
            if start_col >= len(row):
                continue
            cycle_raw = row[start_col]
            discharge_capacity = row[start_col + 2] if start_col + 2 < len(row) else None
            if not (_is_number(cycle_raw) and _is_number(discharge_capacity)):
                continue
            raw_rows.append(
                {
                    "cycle_raw": float(cycle_raw),
                    "charge_capacity_ah": row[start_col + 1] if start_col + 1 < len(row) else np.nan,
                    "discharge_capacity_ah": float(discharge_capacity),
                    "charge_energy_wh": row[start_col + 3] if start_col + 3 < len(row) else np.nan,
                    "discharge_energy_wh": row[start_col + 4] if start_col + 4 < len(row) else np.nan,
                    "capacity_retention_nominal": row[start_col + 5] if start_col + 5 < len(row) else np.nan,
                    "energy_retention_nominal": row[start_col + 6] if start_col + 6 < len(row) else np.nan,
                    "energy_efficiency": row[start_col + 7] if start_col + 7 < len(row) else np.nan,
                }
            )
        if not raw_rows:
            continue

        block_df = pd.DataFrame(raw_rows)
        numeric_cols = [col for col in block_df.columns if col != "cycle_raw"]
        for col in numeric_cols:
            block_df[col] = pd.to_numeric(block_df[col], errors="coerce")

        q_values = block_df["discharge_capacity_ah"].to_numpy(dtype=float)
        peak_idx = int(np.nanargmax(q_values))
        q_bol = float(q_values[peak_idx])
        bol_cycle_raw = float(block_df["cycle_raw"].iloc[peak_idx])

        if block_df["capacity_retention_nominal"].isna().all():
            block_df["capacity_retention_nominal"] = block_df["discharge_capacity_ah"] / nominal_capacity_ah
        if block_df["energy_retention_nominal"].isna().all():
            block_df["energy_retention_nominal"] = block_df["discharge_energy_wh"] / nominal_energy_wh

        block_df["curve_type"] = "measured_cell"
        block_df["cell_id"] = cell_id
        block_df["temperature_c"] = temperature_c
        block_df["rate_p"] = float(rate_p)
        block_df["cycle"] = block_df["cycle_raw"] - bol_cycle_raw
        block_df["capacity_retention_bol"] = block_df["discharge_capacity_ah"] / q_bol
        block_df["q_bol_ah"] = q_bol
        block_df["bol_cycle_raw"] = bol_cycle_raw
        block_df["n_cells"] = 1
        block_df["label"] = _condition_label(temperature_c, rate_p, cell_id)
        records.extend(block_df[DAT_COLUMNS].to_dict("records"))

    if not records:
        return pd.DataFrame(columns=DAT_COLUMNS)
    return pd.DataFrame(records, columns=DAT_COLUMNS)


def _build_mean_curves(measured_df: pd.DataFrame) -> pd.DataFrame:
    if measured_df.empty:
        return pd.DataFrame(columns=DAT_COLUMNS)

    value_cols = [
        "cycle_raw",
        "charge_capacity_ah",
        "discharge_capacity_ah",
        "charge_energy_wh",
        "discharge_energy_wh",
        "capacity_retention_nominal",
        "energy_retention_nominal",
        "energy_efficiency",
        "capacity_retention_bol",
        "q_bol_ah",
        "bol_cycle_raw",
    ]
    grouped = (
        measured_df.groupby(["temperature_c", "rate_p", "cycle"], as_index=False)
        .agg({**{col: "mean" for col in value_cols}, "cell_id": "count"})
        .rename(columns={"cell_id": "n_cells"})
    )
    grouped["curve_type"] = "mean"
    grouped["cell_id"] = "mean"
    grouped["label"] = [
        _condition_label(temp, rate)
        for temp, rate in zip(grouped["temperature_c"], grouped["rate_p"])
    ]
    return grouped[DAT_COLUMNS]


def _prediction_starts(rows: list[list[object]]) -> list[int]:
    starts = []
    for col_idx, value in enumerate(rows[1]):
        if isinstance(value, str) and "预测值" in value:
            starts.append(col_idx)
    return starts


def _prediction_temperature(rows: list[list[object]], start_col: int) -> float:
    preceding = []
    fallback = []
    for col_idx, header in _cell_block_starts(rows):
        try:
            _, temperature_c = _parse_cell_header(header)
        except ValueError:
            continue
        fallback.append((abs(col_idx - start_col), temperature_c))
        if col_idx < start_col:
            preceding.append((col_idx, temperature_c))
    if preceding:
        return max(preceding, key=lambda item: item[0])[1]
    if not fallback:
        raise ValueError(f"Cannot infer prediction temperature near column {start_col + 1}")
    return min(fallback, key=lambda item: item[0])[1]


def _parse_prediction_blocks(rows: list[list[object]], *, rate_p: float) -> pd.DataFrame:
    records = []
    for start_col in _prediction_starts(rows):
        temperature_c = _prediction_temperature(rows, start_col)
        for row in rows[3:]:
            cycle = row[start_col] if start_col < len(row) else None
            retention = row[start_col + 1] if start_col + 1 < len(row) else None
            if not (_is_number(cycle) and _is_number(retention)):
                continue
            records.append(
                {
                    "curve_type": "prediction",
                    "cell_id": "prediction",
                    "label": f"{_condition_label(temperature_c, rate_p)} Prediction",
                    "temperature_c": temperature_c,
                    "rate_p": float(rate_p),
                    "cycle": float(cycle),
                    "cycle_raw": float(cycle),
                    "charge_capacity_ah": np.nan,
                    "discharge_capacity_ah": np.nan,
                    "charge_energy_wh": np.nan,
                    "discharge_energy_wh": np.nan,
                    "capacity_retention_nominal": float(retention),
                    "energy_retention_nominal": np.nan,
                    "energy_efficiency": np.nan,
                    "capacity_retention_bol": float(retention),
                    "q_bol_ah": np.nan,
                    "bol_cycle_raw": np.nan,
                    "n_cells": np.nan,
                }
            )
    if not records:
        return pd.DataFrame(columns=DAT_COLUMNS)
    return pd.DataFrame(records, columns=DAT_COLUMNS)


def parse_cycle_life_excel(
    workbook_path: str | Path,
    *,
    sheet_name: str = "0.25P",
    nominal_capacity_ah: float = 1175.0,
    nominal_energy_wh: float = 3760.0,
    rate_p: float = 0.25,
    include_mean: bool = True,
    include_prediction: bool = True,
) -> pd.DataFrame:
    """Parse a Hithium cycle-life Excel sheet into a normalized long table.

    The returned table includes BOL-aligned cycle numbers. Rows before the
    discharge-capacity peak have negative ``cycle`` values so the raw evidence is
    retained, while downstream loaders can drop those rows for model comparison.
    """
    rows = _read_sheet_rows(workbook_path, sheet_name)
    measured_df = _parse_measured_blocks(
        rows,
        nominal_capacity_ah=float(nominal_capacity_ah),
        nominal_energy_wh=float(nominal_energy_wh),
        rate_p=float(rate_p),
    )
    frames = [measured_df]
    if include_mean:
        frames.append(_build_mean_curves(measured_df))
    if include_prediction:
        frames.append(_parse_prediction_blocks(rows, rate_p=float(rate_p)))
    result = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=DAT_COLUMNS)
    if result.empty:
        return pd.DataFrame(columns=DAT_COLUMNS)
    result = result[DAT_COLUMNS].sort_values(
        ["curve_type", "temperature_c", "cell_id", "cycle"],
        kind="mergesort",
    )
    return result.reset_index(drop=True)


def export_cycle_life_dat(
    workbook_path: str | Path,
    output_path: str | Path,
    **parse_kwargs,
) -> pd.DataFrame:
    """Export a normalized cycle-life table to a tab-separated ``.dat`` file."""
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    df = parse_cycle_life_excel(workbook_path, **parse_kwargs)
    df.to_csv(output, sep="\t", index=False, encoding="utf-8", float_format="%.10g")
    return df


def _normalize_filter_values(values: float | str | Iterable[float | str] | None) -> set[float] | None:
    if values is None:
        return None
    if isinstance(values, (str, int, float)):
        values = [values]
    normalized = set()
    for value in values:
        if isinstance(value, str):
            match = re.search(r"-?\d+(?:\.\d+)?", value)
            if not match:
                continue
            normalized.add(float(match.group(0)))
        else:
            normalized.add(float(value))
    return normalized


def _filter_float_column(df: pd.DataFrame, column: str, allowed: set[float] | None) -> pd.DataFrame:
    if allowed is None:
        return df
    return df[df[column].astype(float).round(10).isin({round(value, 10) for value in allowed})]


def load_lifecycle_dat(
    file_path: str | Path,
    *,
    curve_type: str | Iterable[str] = "mean",
    retention_col: str = "capacity_retention_bol",
    drop_pre_bol: bool = True,
    temperature_c: float | str | Iterable[float | str] | None = None,
    rate_p: float | str | Iterable[float | str] | None = None,
) -> list[dict]:
    """Load normalized lifecycle ``.dat`` data as ``compare_all`` experiment dicts."""
    df = pd.read_csv(file_path, sep="\t")
    missing = [col for col in DAT_COLUMNS if col not in df.columns]
    if missing:
        raise ValueError(f"Lifecycle dat is missing required columns: {missing}")
    if retention_col not in df.columns:
        raise ValueError(f"Retention column {retention_col!r} not found")

    curve_types = {curve_type} if isinstance(curve_type, str) else set(curve_type)
    data = df[df["curve_type"].isin(curve_types)].copy()
    data = _filter_float_column(data, "temperature_c", _normalize_filter_values(temperature_c))
    data = _filter_float_column(data, "rate_p", _normalize_filter_values(rate_p))
    if drop_pre_bol:
        data = data[data["cycle"] >= 0]

    exp_data = []
    group_cols = ["curve_type", "temperature_c", "rate_p", "cell_id", "label"]
    for (_, _, _, _, label), group in data.groupby(group_cols, sort=True, dropna=False):
        group = group.sort_values("cycle", kind="mergesort")
        retention = pd.to_numeric(group[retention_col], errors="coerce").to_numpy(dtype=float)
        exp_data.append(
            {
                "label": str(label),
                "cycle": pd.to_numeric(group["cycle"], errors="coerce").to_numpy(dtype=float),
                "discharge_capacity": pd.to_numeric(group["discharge_capacity_ah"], errors="coerce").to_numpy(dtype=float),
                "charge_capacity": pd.to_numeric(group["charge_capacity_ah"], errors="coerce").to_numpy(dtype=float),
                "retention": retention,
                "retention_bol": pd.to_numeric(group["capacity_retention_bol"], errors="coerce").to_numpy(dtype=float),
                "retention_nominal": pd.to_numeric(group["capacity_retention_nominal"], errors="coerce").to_numpy(dtype=float),
                "efficiency": pd.to_numeric(group["energy_efficiency"], errors="coerce").to_numpy(dtype=float),
                "discharge_energy": pd.to_numeric(group["discharge_energy_wh"], errors="coerce").to_numpy(dtype=float),
                "charge_energy": pd.to_numeric(group["charge_energy_wh"], errors="coerce").to_numpy(dtype=float),
                "temperature": np.full(len(group), float(group["temperature_c"].iloc[0])),
                "max_force": np.array([]),
                "min_force": np.array([]),
                "source_file": str(file_path),
                "curve_type": str(group["curve_type"].iloc[0]),
            }
        )
    return exp_data
