"""Cycle metrics / summary 报表导出 (Excel)。

此前位于 ``analysis.py``，因属于纯报表写出（pandas + openpyxl），
与分析计算解耦后独立成模块。
"""
import logging
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .analysis import compute_cycle_energies, extract_all_metrics_from_sol, retention_from_capacity

logger = logging.getLogger(__name__)


def export_cycle_metrics_report(
    sol_list,
    sim_labels_list,
    cycle_step=50,
    output_filename="Cycle_Metrics_Report.xlsx",
):
    """按每个工况四列分组导出循环结果到 Excel。

    参数
    ----
    sol_list : list
        仿真结果列表。
    sim_labels_list : list[str]
        仿真标签列表，与 sol_list 一一对应。
    cycle_step : int
        每个仿真循环对应的真实循环步长。
    output_filename : str or Path or None
        Excel 输出路径；设为 None 时仅返回 DataFrame，不写文件。

    返回
    ----
    pandas.DataFrame or None
        汇总后的宽表，列名格式为 "{label}_{metric}"。
    """
    metric_headers = ["Cycle", "SOH (%)", "Discharge Capacity (Ah)", "Efficiency (%)"]
    thin_side = Side(style="thin", color="D9D9D9")
    header_fill = PatternFill(fill_type="solid", fgColor="D9D9D9")
    group_fill = PatternFill(fill_type="solid", fgColor="BFBFBF")
    center_alignment = Alignment(horizontal="center", vertical="center")

    condition_tables = []
    summary_blocks = []

    for idx, (sol, label) in enumerate(zip(sol_list, sim_labels_list), start=1):
        res = compute_cycle_energies(sol)
        caps = np.asarray(res.get("discharge_cap", np.array([])), dtype=float)
        effs = np.asarray(res.get("efficiency", np.array([])), dtype=float) * 100

        valid_mask = np.isfinite(caps) & np.isfinite(effs)
        caps = caps[valid_mask]
        effs = effs[valid_mask]

        if caps.size == 0:
            logger.warning("skip %s because no valid cycle data was found", label)
            continue

        cycles = np.arange(1, caps.size + 1, dtype=int) * int(cycle_step)
        soh = retention_from_capacity(caps) * 100
        clean_label = str(label).strip() or f"Condition {idx}"

        df_single = pd.DataFrame(
            {
                "Cycle": cycles,
                "SOH (%)": soh,
                "Discharge Capacity (Ah)": caps,
                "Efficiency (%)": effs,
            }
        )
        condition_tables.append((clean_label, df_single))
        summary_blocks.append(
            df_single.rename(
                columns={
                    "Cycle": f"{clean_label}_Cycle",
                    "SOH (%)": f"{clean_label}_SOH(%)",
                    "Discharge Capacity (Ah)": f"{clean_label}_DischargeCapacity(Ah)",
                    "Efficiency (%)": f"{clean_label}_Efficiency(%)",
                }
            )
        )

    if not condition_tables:
        logger.warning("No cycle metrics were generated for export.")
        return None

    full_df = pd.concat(summary_blocks, axis=1)
    if output_filename is None:
        return full_df

    output_path = Path(output_filename)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Cycle Metrics"

    for block_idx, (label, df_single) in enumerate(condition_tables):
        start_col = block_idx * 4 + 1
        end_col = start_col + 3

        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        group_cell = ws.cell(row=1, column=start_col, value=label)
        group_cell.fill = group_fill
        group_cell.font = Font(bold=True)
        group_cell.alignment = center_alignment

        for col_offset, header in enumerate(metric_headers):
            cell = ws.cell(row=2, column=start_col + col_offset, value=header)
            cell.fill = header_fill
            cell.font = Font(bold=True)
            cell.alignment = center_alignment
            cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        for row_offset, row_values in enumerate(df_single.itertuples(index=False, name=None), start=3):
            for col_offset, value in enumerate(row_values):
                cell = ws.cell(row=row_offset, column=start_col + col_offset, value=float(value) if col_offset else int(value))
                cell.alignment = center_alignment
                cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        data_end_row = len(df_single) + 2
        for metric_col in range(start_col + 1, end_col + 1):
            if data_end_row >= 3:
                col_letter = ws.cell(row=2, column=metric_col).column_letter
                ws.conditional_formatting.add(
                    f"{col_letter}3:{col_letter}{data_end_row}",
                    ColorScaleRule(
                        start_type="min",
                        start_color="F8696B",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFEB84",
                        end_type="max",
                        end_color="63BE7B",
                    ),
                )

    for column_index in range(1, ws.max_column + 1):
        max_length = 0
        column_letter = ws.cell(row=2, column=column_index).column_letter
        for row_index in range(1, ws.max_row + 1):
            cell = ws.cell(row=row_index, column=column_index)
            cell_value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(cell_value))
            cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
            if cell.row == 1:
                cell.fill = group_fill
                cell.font = Font(bold=True)
                cell.alignment = center_alignment
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 14), 28)

    ws.freeze_panes = "B3"
    wb.save(output_path)
    logger.info("Exported cycle metrics to: %s", output_path)
    return full_df


def export_full_metrics_summary(sol_list, sim_labels_list, cycle_step=50, output_filename="Full_Metrics_Summary.xlsx"):
    """导出多工况的圈数/容量/保持率/能量/能效汇总 Excel。"""
    all_dfs = []
    for sol, label in zip(sol_list, sim_labels_list):
        caps, e_chgs, e_dchgs, effs = extract_all_metrics_from_sol(sol)
        if len(caps) == 0:
            logger.warning("跳过 %s: 无有效数据", label)
            continue

        cycles = np.arange(1, len(caps) + 1) * cycle_step
        retention = retention_from_capacity(caps)

        df_single = pd.DataFrame(
            {
                f"{label}_Cycle": cycles,
                f"{label}_Capacity(Ah)": caps,
                f"{label}_Retention(%)": retention * 100,
                f"{label}_Energy_Charge(Wh)": e_chgs,
                f"{label}_Energy_Discharge(Wh)": e_dchgs,
                f"{label}_Efficiency(%)": effs * 100,
            }
        )
        all_dfs.append(df_single)

    if not all_dfs:
        logger.warning("未生成任何数据。")
        return None

    full_df = pd.concat(all_dfs, axis=1)
    if output_filename is not None:
        full_df.to_excel(output_filename, index=False)
        logger.info("成功导出全量汇总数据至: %s", output_filename)
    return full_df
