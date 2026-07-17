import numpy as np
import openpyxl

from src.lifecycle_exp import export_cycle_life_dat, load_lifecycle_dat, parse_cycle_life_excel


def _write_cell_block(ws, start_col, label, rows):
    ws.cell(1, start_col, label)
    headers = [
        "圈数",
        "充电容量 (Ah)",
        "放电容量 (Ah)",
        "充电能量 (Wh)",
        "放电能量 (Wh)",
        "标称容量保持率",
        "标称能量保持率",
        "能量效率",
    ]
    for offset, header in enumerate(headers):
        ws.cell(2, start_col + offset, header)
    for row_idx, values in enumerate(rows, start=4):
        for offset, value in enumerate(values):
            ws.cell(row_idx, start_col + offset, value)


def _build_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "0.25P"
    _write_cell_block(
        ws,
        1,
        "Cell 01 25°C",
        [
            [1, 101.0, 100.0, 330.0, 320.0, None, None, 0.970],
            [2, 106.0, 105.0, 340.0, 330.0, None, None, 0.971],
            [3, 103.0, 102.0, 335.0, 325.0, None, None, 0.970],
        ],
    )
    _write_cell_block(
        ws,
        10,
        "Cell 02 25°C",
        [
            [1, 100.0, 99.0, 330.0, 319.0, None, None, 0.967],
            [2, 105.0, 104.0, 339.0, 329.0, None, None, 0.970],
            [3, 102.0, 101.0, 334.0, 324.0, None, None, 0.970],
        ],
    )
    _write_cell_block(
        ws,
        22,
        "Cell 01 45°C",
        [
            [1, 111.0, 110.0, 350.0, 340.0, None, None, 0.971],
            [2, 109.0, 108.0, 346.0, 336.0, None, None, 0.971],
        ],
    )
    _write_cell_block(
        ws,
        31,
        "Cell 02 45°C",
        [
            [1, 112.0, 111.0, 351.0, 341.0, None, None, 0.972],
            [2, 110.0, 109.0, 347.0, 337.0, None, None, 0.971],
        ],
    )
    ws.cell(2, 19, "标称容量保持率-预测值")
    ws.cell(4, 19, 10)
    ws.cell(4, 20, 0.99)
    ws.cell(2, 40, "标称容量保持率-预测值")
    ws.cell(4, 40, 10)
    ws.cell(4, 41, 0.98)
    wb.save(path)


def test_parse_cycle_life_excel_builds_bol_aligned_long_table(tmp_path):
    workbook = tmp_path / "cycle.xlsx"
    _build_workbook(workbook)

    df = parse_cycle_life_excel(workbook, nominal_capacity_ah=100.0, nominal_energy_wh=320.0)

    measured = df[df["curve_type"] == "measured_cell"]
    mean = df[df["curve_type"] == "mean"]
    prediction = df[df["curve_type"] == "prediction"]
    assert measured["cell_id"].nunique() == 2
    assert len(mean[mean["temperature_c"] == 25.0]) == 3
    assert len(mean[mean["temperature_c"] == 45.0]) == 2
    assert len(prediction) == 2

    cell_01_25 = measured[(measured["cell_id"] == "Cell 01") & (measured["temperature_c"] == 25.0)]
    assert cell_01_25["cycle"].tolist() == [-1.0, 0.0, 1.0]
    assert np.isclose(cell_01_25.loc[cell_01_25["cycle"] == 0, "capacity_retention_bol"].iloc[0], 1.0)

    mean_25_last = mean[(mean["temperature_c"] == 25.0) & (mean["cycle"] == 1.0)].iloc[0]
    expected = np.mean([102.0 / 105.0, 101.0 / 104.0])
    assert np.isclose(mean_25_last["capacity_retention_bol"], expected)
    assert mean_25_last["n_cells"] == 2


def test_export_and_load_lifecycle_dat_round_trip(tmp_path):
    workbook = tmp_path / "cycle.xlsx"
    output = tmp_path / "cycle.dat"
    _build_workbook(workbook)

    export_cycle_life_dat(workbook, output, nominal_capacity_ah=100.0, nominal_energy_wh=320.0)
    exp_data = load_lifecycle_dat(output, curve_type="mean")

    labels = {item["label"] for item in exp_data}
    assert labels == {"25°C 0.25P", "45°C 0.25P"}
    exp_25 = next(item for item in exp_data if item["label"] == "25°C 0.25P")
    assert exp_25["cycle"].tolist() == [0.0, 1.0]
    assert np.isclose(exp_25["retention"][0], 1.0)
    assert np.isclose(exp_25["retention"][1], np.mean([102.0 / 105.0, 101.0 / 104.0]))

